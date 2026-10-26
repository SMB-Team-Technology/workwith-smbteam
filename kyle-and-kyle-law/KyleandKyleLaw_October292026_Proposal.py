"""
Audit PowerPoint Template — SMB Team
======================================
Generates a 3-slide proposal PPTX from the Growth, Profit, and Freedom Roadmap.
Uses python-pptx. Do not modify the layout engine below the FILL section.
Only replace the # FILL: placeholders with audit-specific content.

Slide 1 — Where [Firm] Stands Today         (what you told us, stage, key gaps)
Slide 2 — Your Growth Plan: 3 Priorities    (action plan)
Slide 3 — Your Investment & What's Next     (pricing + how this plan addresses it)

Output: [friendly-name]/[FirmName]_[Date]_Proposal.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from functools import lru_cache
from PIL import ImageFont
import os
import re

# ═══════════════════════════════════════════════════════════════════
# FILL — Replace every placeholder. Do not delete any variable.
# ═══════════════════════════════════════════════════════════════════

# Report metadata
FIRM_NAME      = "Kyle and Kyle Law"
SALES_REP      = "Jonathan Farace"
OUTPUT_PATH    = "kyle-and-kyle-law/KyleandKyleLaw_October292026_Proposal.pptx"

PROSPECT_QUOTES = [
    ("You said", "Emails is a huge one for me… we have one person hired to review three people's emails"),
    ("You said", "Just the running the business pulls me away from it."),
    ("You said", "It requires too much double work because then we still have to put it in our case management system"),
]

KEY_GAPS = [
    ("Team", "Owners still run the business, which keeps pulling them away from growth"),
    ("Intake", "Double entry before each case opens costs staff time on every new client"),
    ("Profit Plan", "No current goal or fee-recovery tracking makes expansion hard to plan"),
]

STAGE_TEXT = "Stage 4: Small Business Manager  →  Goal: Stage 6, Law Firm Owner"

SLIDE_2_TITLE = "Your Growth Plan: 3 Priorities to Scale Into New States"

PRIORITIES = [
    (
        "Put AI to Work", "On Routine Tasks", "003A59",
        [
            "Monthly 1:1 with your Fractional CTO sets the roadmap",
            "Automate email triage and calendaring first",
            "Remove double entry at onboarding",
            "Build the Loan Mod Analyzer for faster file review",
            "Turn routine work into time back for you",
        ],
    ),
    (
        "Equip the Team", "To Use AI Safely", "69CD2B",
        [
            "Give all 15 team members AI Workforce Pro access",
            "Weekly implementation session keeps rollout on track",
            "Set an AI policy and data-access rules",
            "Protect client data while the team moves faster",
            "One setup, not a pile of disconnected tools",
        ],
    ),
    (
        "Free Your Time", "To Scale New States", "003A59",
        [
            "Step back from running the business day to day",
            "Track fee-recovery revenue with your dashboards",
            "Set a fresh revenue target for the expansion",
            "Make Indiana and Ohio repeatable, not heavier",
            "Add ops and finance leadership as you grow",
        ],
    ),
]

PACKAGES = [
    (
        "FRACTIONAL CTO — LEVEL 1",
        "$3,297", "$3,797/mo",
        "Monthly 1:1 · AI skills · Custom automation · Reports",
        "0091C9",
    ),
    (
        "AI WORKFORCE PRO — 15 SEATS",
        "$2,797", "",
        "Team-wide AI access · Weekly implementation session",
        "003A59",
    ),
]

BUNDLE_TOTAL   = "$6,094 / mo"
BUNDLE_SAVINGS = "Save $500/mo by bundling"

AD_SPEND_NOTE = None

AVG_CASE_VALUE     = "$X,XXX"
CONSERVATIVE_LABEL = "Conservative  (X cases/mo):"
CONSERVATIVE_RESULT = "$XX,XXX revenue · X.X× ROAS"
AGGRESSIVE_LABEL   = "Aggressive  (X cases/mo):"
AGGRESSIVE_RESULT  = "$XX,XXX revenue · X.X× ROAS"

WHAT_YOU_TOLD_US = [
    ("AI roadmap", "Monthly 1:1 builds your roadmap; report tracks wins"),
    ("Email overload", "Up to 2 skills + 1 automation a month start there"),
    ("Tool sprawl", "Whole team in one place, weekly sessions"),
]

CLOSING_QUOTE = (
    '"The clients are already coming — the AI leverage to stop running the business '
    'and start scaling it is the missing piece."'
)

# ═══════════════════════════════════════════════════════════════════
# LAYOUT ENGINE — DO NOT MODIFY BELOW THIS LINE
# ═══════════════════════════════════════════════════════════════════

# ── Color palette ─────────────────────────────────────────────────
def rgb(h): return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

# SMB Team brand colors (see brand book): Deep Wood Blue, Ocean Blue, Lime Green.
# These three anchor the palette and are unchanged from the previous design —
# what changed in this pass is how they're applied (thin accents/pills instead
# of large solid fills) plus several supporting tones below, retuned to softer,
# lower-saturation values matched from the reference redesign.
NAVY        = rgb("003A59")   # Deep Wood Blue — primary brand color
OCEAN_BLUE  = rgb("0091C9")   # brand accent — large numbers, decorative accents
LIME_GREEN  = rgb("69CD2B")   # brand accent — small "pop" text/labels on navy
WHITE       = rgb("FFFFFF")

BG_WASH       = rgb("F4F7F9")   # full-bleed slide background
URGENCY_CARD  = rgb("0B4C6E")   # muted navy — staircase steps not yet reached
URGENCY_SUB   = rgb("9CC0D6")   # step numbers on muted staircase steps

GREEN       = rgb("3F8F1A")

DARK_TEXT   = rgb("1B2A38")   # near-black body/label text on white cards
BODY_TEXT   = rgb("4A5C6B")   # neutral gray-blue for card copy (key gaps, ad spend note, plan rows)
SLATE       = rgb("7C8D9B")   # lighter captions (retail price, services line)
FOOTER_TEXT = rgb("8FB2C6")

ROI_BG      = rgb("F0FAF4")
ROI_GREEN   = rgb("0B6B4A")
ROI_HILIGHT = rgb("E2F3E9")

BUNDLE_SAVINGS_TEXT = rgb("A9D68C")   # lighter green — reads as "money saved"
GOAL_TEXT_DARK      = rgb("143306")   # dark green text on Lime Green fills

# Slide 2 priority columns always alternate navy / lime / navy, whatever hex
# the FILL section passes — no Ocean Blue or other accents on that slide.
# Each entry: (header fill, title text, "PRIORITY N" kicker text). Lime
# Green is light, so its header uses dark green text instead of white.
PRIORITY_STYLE = [
    (NAVY,       WHITE,          LIME_GREEN),
    (LIME_GREEN, GOAL_TEXT_DARK, GOAL_TEXT_DARK),
    (NAVY,       WHITE,          LIME_GREEN),
]

FONT = "Poppins"
LOGO_PATH = os.path.join(os.path.dirname(__file__), "smb_team_logo.png")
# This script is copied into each firm's own folder (a sibling of "Design
# Files/" at the repo root) to run — the logo above sits right next to it
# because it's copied too, but fonts/ is not, so it's addressed relative to
# the repo root instead.
FONT_DIR = os.path.join(os.path.dirname(__file__), "..", "Design Files", "fonts")

# ── Core helpers ──────────────────────────────────────────────────

def emu(*inches):
    return tuple(int(x * 914400) for x in inches)


# Fields with a documented character budget (see .claude/commands/audit-pptx.md)
# get hard-capped where they're placed on the slide — a backstop so oversized
# content can never spill past its box, no matter what the FILL section
# contains. auto-shrink (set in add_text below) is the first line of defense;
# this cap keeps the shrunk font from having to go microscopic.
_TRUNCATED = []  # (label, original_text) pairs — reported just before saving

def cap(text, max_chars, label=""):
    text = "" if text is None else str(text)
    if len(text) <= max_chars:
        return text
    truncated = text[:max_chars - 1].rstrip()
    if " " in truncated:
        truncated = truncated.rsplit(" ", 1)[0]
    truncated = truncated.rstrip(" ,.;:-—–") + "…"
    _TRUNCATED.append((label or text[:24], text))
    return truncated


def add_rect(slide, left, top, w, h, fill=None, line=False):
    from pptx.util import Emu as E
    shape = slide.shapes.add_shape(1, E(int(left*914400)), E(int(top*914400)),
                                   E(int(w*914400)), E(int(h*914400)))
    shape.line.fill.background() if not line else None
    if not line:
        shape.line.color.rgb = WHITE
        shape.line.width = 0
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    return shape


# ── Text fitting ──────────────────────────────────────────────────
# PowerPoint's own "shrink text on overflow" (normAutofit) is only a flag:
# python-pptx can't compute the scale, PowerPoint applies it only after a
# human edits the box, and Google Slides / Keynote / previews ignore it. So
# every text box is measured here instead, against the same Poppins files
# that get embedded, and the font size actually written to the file is the
# one that fits — every viewer then renders exactly the same thing.

_INSET_X = 0.10        # python-pptx's default left/right text-box padding (in)
_INSET_Y = 0.05        # ... and top/bottom padding (in)
_WIDTH_SAFETY = 0.97   # small margin for renderer kerning differences
_FIT_STEP = 0.5        # "drop one font size" = 0.5pt per step
_SHRUNK = []           # (label, from_pt, to_pt) — reported just before saving


@lru_cache(maxsize=None)
def _font_file(bold, italic):
    name = "Poppins-Bold.ttf" if bold else ("Poppins-Italic.ttf" if italic else "Poppins-Regular.ttf")
    # Measured at 1000px; widths scale linearly with point size.
    return ImageFont.truetype(os.path.join(FONT_DIR, name), 1000)


def text_width(text, size, bold=False, italic=False):
    """Rendered width of a single line of text, in inches."""
    return _font_file(bold, italic).getlength(text) * size / 1000 / 72


def line_height(size, bold=False, italic=False):
    """Single-spaced line height in inches (Poppins: ascent + descent ≈ 1.4em)."""
    ascent, descent = _font_file(bold, italic).getmetrics()
    return (ascent + descent) * size / 1000 / 72


def baseline_offset(size, bold=False, italic=False):
    """Distance from a text box's top edge to its first line's baseline."""
    ascent, _ = _font_file(bold, italic).getmetrics()
    return _INSET_Y + ascent * size / 1000 / 72


def wrap_lines(text, size, max_w, bold=False, italic=False):
    """Greedy word wrap the way PowerPoint does it. Returns the lines."""
    lines = []
    for para in str(text).split("\n"):
        cur = ""
        for word in para.split():
            trial = f"{cur} {word}" if cur else word
            if cur and text_width(trial, size, bold, italic) > max_w:
                lines.append(cur)
                cur = word
            else:
                cur = trial
        lines.append(cur)
    return lines


def fit_text(text, w, h, size, bold=False, italic=False, max_lines=None, min_size=None):
    """Largest size <= `size` (in _FIT_STEP steps, down to `min_size`) at which
    `text` wraps inside a w x h box. Returns (size, lines), or (None, None)
    when it doesn't fit even at min_size. A box is always allowed one line
    even if it's shorter than that line — small single-line labels are
    positioned for that on purpose; what must never happen is extra lines."""
    if min_size is None:
        min_size = max(6.0, round(size * 0.75 * 2) / 2)
    max_w = (w - 2 * _INSET_X) * _WIDTH_SAFETY
    avail_h = h - 2 * _INSET_Y
    s = size
    while s >= min_size - 1e-6:
        lines = wrap_lines(text, s, max_w, bold, italic)
        fits_h = max(1, int(avail_h / line_height(s, bold, italic) + 0.15))
        allowed = min(fits_h, max_lines) if max_lines else fits_h
        if len(lines) <= allowed and all(text_width(l, s, bold, italic) <= max_w for l in lines):
            return s, lines
        s -= _FIT_STEP
    return None, None


def _truncate_to_fit(text, w, h, size, bold, italic, max_lines):
    words = str(text).split()
    while len(words) > 1:
        words.pop()
        trial = " ".join(words).rstrip(" ,.;:-—–·") + "…"
        if fit_text(trial, w, h, size, bold, italic, max_lines, min_size=size)[0]:
            return trial
    return words[0] if words else ""


def add_text(slide, text, left, top, w, h, size, color, bold=False,
             align=PP_ALIGN.LEFT, italic=False, wrap=True,
             cap_chars=None, cap_label="", fit=True, max_lines=None,
             min_size=None, anchor=None):
    from pptx.util import Emu as E, Pt as P
    if cap_chars is not None:
        text = cap(text, cap_chars, label=cap_label)
    if fit and text:
        fitted, _ = fit_text(text, w, h, size, bold, italic, max_lines, min_size)
        if fitted is None:
            # Doesn't fit even at the smallest allowed size — keep the floor
            # size and trim at a word boundary instead of spilling the box.
            floor = min_size if min_size is not None else max(6.0, round(size * 0.75 * 2) / 2)
            original = text
            text = _truncate_to_fit(text, w, h, floor, bold, italic, max_lines)
            _TRUNCATED.append((cap_label or original[:24], original))
            fitted = floor
        if fitted < size:
            _SHRUNK.append((cap_label or str(text)[:32], size, fitted))
        size = fitted
    txb = slide.shapes.add_textbox(E(int(left*914400)), E(int(top*914400)),
                                   E(int(w*914400)), E(int(h*914400)))
    tf = txb.text_frame
    tf.word_wrap = wrap
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = P(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    # The size above is already the measured fit — turn PowerPoint's own
    # autofit off so no viewer rescales it a second time.
    tf.auto_size = MSO_AUTO_SIZE.NONE
    return txb


def add_title(slide, text, w):
    """Slide title inside the 0.92in navy header. One line at 20pt, dropping
    a half-size at a time to 15pt; if a long firm name still doesn't fit,
    two bottom-anchored lines from 15pt down to 12pt. Never crosses the
    header's bottom edge at 0.92in."""
    one_line, _ = fit_text(text, w, 0.44, 20, bold=True, max_lines=1, min_size=15)
    if one_line:
        add_text(slide, text, 0.32, 0.42, w, 0.44, one_line, WHITE, bold=True,
                 fit=False, anchor=MSO_ANCHOR.MIDDLE)
    else:
        add_text(slide, text, 0.32, 0.30, w, 0.60, 15, WHITE, bold=True,
                 max_lines=2, min_size=12, anchor=MSO_ANCHOR.BOTTOM,
                 cap_label="slide title")


def add_header(slide, title, title_w):
    """Background wash, navy header, kicker and title — shared by all three
    slides. No divider stripe under the header: the deck stays quiet so it
    supports the rep instead of competing."""
    add_rect(slide, 0, 0, 10, 5.625, fill=BG_WASH)
    add_rect(slide, 0, 0, 10, 0.92, fill=NAVY)
    add_text(slide, f"LAW FIRM GROWTH AUDIT  ·  {FIRM_NAME.upper()}",
             0.32, 0.20, title_w, 0.18, 7.5, LIME_GREEN, bold=True,
             max_lines=1, cap_label="header kicker")
    add_title(slide, title, title_w)


def add_footer(slide, page, total, logo_path):
    add_rect(slide, 0, 5.27, 10, 0.35, fill=NAVY)
    if logo_path and os.path.isfile(logo_path):
        from pptx.util import Emu as E
        slide.shapes.add_picture(logo_path,
                                 E(int(0.30*914400)), E(int(5.33*914400)),
                                 E(int(1.03*914400)), E(int(0.21*914400)))
    add_text(slide, f"CONFIDENTIAL  ·  PREPARED BY {SALES_REP.upper()} | SMB TEAM",
             1.55, 5.35, 5.40, 0.20, 7, FOOTER_TEXT)
    add_text(slide, f"{page} / {total}", 8.40, 5.35, 1.30, 0.20, 7, FOOTER_TEXT,
             align=PP_ALIGN.RIGHT)


_BULLET_INDENT = 0.16   # hanging indent for bullet text (in)
_BULLET_GAP_PT = 6      # space before each bullet after the first


def add_bullets(slide, items, left, top, w, h, size, color, bullet_color, label=""):
    """One text box with a real bulleted paragraph per item, sized so all
    items fit: drops 0.5pt at a time (down to 75%) until the wrapped lines
    plus inter-bullet spacing fit the box."""
    from pptx.util import Emu as E, Pt as P
    max_w = (w - 2 * _INSET_X - _BULLET_INDENT) * _WIDTH_SAFETY
    avail_h = h - 2 * _INSET_Y
    floor = max(6.0, round(size * 0.75 * 2) / 2)
    s = size
    while s > floor:
        n_lines = sum(len(wrap_lines(t, s, max_w)) for t in items)
        need = n_lines * line_height(s) + (len(items) - 1) * _BULLET_GAP_PT / 72
        if need <= avail_h:
            break
        s -= _FIT_STEP
    if s < size:
        _SHRUNK.append((label, size, s))

    txb = slide.shapes.add_textbox(E(int(left*914400)), E(int(top*914400)),
                                   E(int(w*914400)), E(int(h*914400)))
    tf = txb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    hex_bullet = "%02X%02X%02X" % tuple(bullet_color)
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i:
            p.space_before = P(_BULLET_GAP_PT)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(_BULLET_INDENT * 914400)))
        pPr.set("indent", str(-int(_BULLET_INDENT * 914400)))
        for tag, attrs in [("a:buClr", None), ("a:buFont", {"typeface": "Arial"}),
                           ("a:buChar", {"char": "•"})]:
            el = etree.SubElement(pPr, qn(tag))
            if tag == "a:buClr":
                etree.SubElement(el, qn("a:srgbClr")).set("val", hex_bullet)
            else:
                for k, v in attrs.items():
                    el.set(k, v)
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = P(s)
        run.font.color.rgb = color
    return txb


def parse_stage_text(text):
    """Split STAGE_TEXT ("Stage 4: Small Business Manager → Goal: Stage 6,
    Law Firm Owner") into (current, goal, current_n, goal_n). Any part that
    can't be found comes back as None."""
    parts = [p.strip() for p in re.split(r"\s*(?:→|->)\s*", str(text), maxsplit=1)]
    current = parts[0] or None
    goal = parts[1] if len(parts) > 1 and parts[1] else None
    def stage_n(s):
        m = re.search(r"Stage\s*(\d)", s or "", re.I)
        n = int(m.group(1)) if m else None
        return n if n and 1 <= n <= 6 else None
    return current, goal, stage_n(current), stage_n(goal)


# ── Slide 1: Where You Stand Today ────────────────────────────────

# Legacy staircase — stage number → short name (audit-write.md "FIRM STAGE").
STAGE_NAMES = {2: "The Lawyer", 3: "Solo Practitioner", 4: "Small Business Manager",
               5: "Law Firm CEO", 6: "Law Firm Owner"}


def add_rich_text(slide, paras, left, top, w, h, label="", anchor=None, min_scale=0.75):
    """One text box holding several paragraphs, each a list of runs
    (text, size, bold, italic, color) plus a space-before in pt. Every size
    is scaled down together, a half-point at a time on the largest run,
    until all wrapped lines fit the box."""
    from pptx.util import Emu as E, Pt as P
    max_w = (w - 2 * _INSET_X) * _WIDTH_SAFETY
    avail_h = h - 2 * _INSET_Y
    top_size = max(r[1] for runs, _ in paras for r in runs)
    scale = 1.0
    while True:
        need = 0.0
        for i, (runs, space) in enumerate(paras):
            # Measure a paragraph at its largest run's style — runs in one
            # paragraph here share a line style (label vs quote).
            big = max(runs, key=lambda r: r[1])
            s = big[1] * scale
            text = "".join(r[0] for r in runs)
            need += len(wrap_lines(text, s, max_w, big[2], big[3])) * line_height(s)
            if i:
                need += space * scale / 72
        if need <= avail_h or scale <= min_scale:
            break
        scale = max(min_scale, (top_size * scale - _FIT_STEP) / top_size)
    if scale < 1.0:
        _SHRUNK.append((label, top_size, round(top_size * scale * 2) / 2))

    txb = slide.shapes.add_textbox(E(int(left*914400)), E(int(top*914400)),
                                   E(int(w*914400)), E(int(h*914400)))
    tf = txb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    if anchor is not None:
        tf.vertical_anchor = anchor
    for i, (runs, space) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i and space:
            p.space_before = P(space * scale)
        for text, size, bold, italic, color in runs:
            r = p.add_run()
            r.text = text
            r.font.name = FONT
            r.font.size = P(round(size * scale * 2) / 2)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
    return txb


def add_rule(slide, x1, y1, x2, y2, color, weight_pt=0.75):
    """A thin straight divider line (a connector, so no fill or shadow)."""
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Emu as E, Pt as P
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(int(x1*914400)),
                                      E(int(y1*914400)), E(int(x2*914400)), E(int(y2*914400)))
    conn.line.color.rgb = color
    conn.line.width = P(weight_pt)


def build_staircase_card(slide, x, y, w, h):
    """Navy card: where the firm is today and where it needs to be, on the
    legacy staircase (stages 2–6). Current step filled lime, goal step
    outlined lime, steps already climbed muted green, the rest muted navy."""
    from pptx.util import Emu as E, Pt as P
    add_rect(slide, x, y, w, h, fill=NAVY)
    current, goal, cur_n, goal_n = parse_stage_text(STAGE_TEXT)

    def stage_label(text, n):
        # The official staircase name for the stage number, so it always
        # matches the audit; free text only if no stage number was given.
        if n in STAGE_NAMES:
            return STAGE_NAMES[n]
        return re.sub(r"^\s*(?:Goal:\s*)?Stage\s*\d\s*[:,–—-]?\s*", "", text or "", flags=re.I)

    half = (w - 0.50) / 2
    cur_name = stage_label(current, cur_n)
    goal_name = stage_label(goal, goal_n) if goal else ""
    # Both stage names share one size so the two ends read as a pair.
    name_size = min(fit_text(t, half, 0.46, 12, bold=True, max_lines=2, min_size=9)[0] or 9
                    for t in [cur_name, goal_name] if t) if (cur_name or goal_name) else 12
    cur_kicker = f"TODAY  ·  STAGE {cur_n}" if cur_n else "TODAY"
    add_text(slide, cur_kicker, x+0.20, y+0.12, half, 0.18, 7.5, LIME_GREEN, bold=True,
             fit=False)
    add_text(slide, cur_name, x+0.20, y+0.28, half, 0.46, name_size, WHITE, bold=True,
             max_lines=2, min_size=name_size, cap_label="STAGE_TEXT current")
    if goal:
        goal_kicker = f"GOAL  ·  STAGE {goal_n}" if goal_n else "GOAL"
        add_text(slide, goal_kicker, x+0.30+half, y+0.12, half, 0.18, 7.5,
                 LIME_GREEN, bold=True, fit=False, align=PP_ALIGN.RIGHT)
        add_text(slide, goal_name, x+0.30+half, y+0.28, half, 0.46, name_size,
                 WHITE, bold=True, max_lines=2, min_size=name_size, align=PP_ALIGN.RIGHT,
                 cap_label="STAGE_TEXT goal")

    if cur_n is None:
        return
    stages = [2, 3, 4, 5, 6]
    gap = 0.06
    # Same staircase size whether the card is half or full width; on a
    # full-width card it sits to the right, under the goal label.
    area_w = min(w - 0.40, 3.80)
    area_x = x + w - 0.20 - area_w
    step_w = (area_w - gap * (len(stages) - 1)) / len(stages)
    base = y + h - 0.16
    top_room = y + 0.86
    unit = (base - top_room) / len(stages)
    for i, n in enumerate(stages):
        sx = area_x + i * (step_w + gap)
        sh_h = unit * (i + 1)
        bar = slide.shapes.add_shape(1, E(int(sx*914400)), E(int((base-sh_h)*914400)),
                                     E(int(step_w*914400)), E(int(sh_h*914400)))
        bar.fill.solid()
        if n == cur_n:
            bar.fill.fore_color.rgb = LIME_GREEN
            bar.line.fill.background()
            num_color = GOAL_TEXT_DARK
        elif n < cur_n:
            bar.fill.fore_color.rgb = rgb("2F6B3A")
            bar.line.fill.background()
            num_color = rgb("A9D68C")
        elif n == goal_n:
            bar.fill.fore_color.rgb = NAVY
            bar.line.color.rgb = LIME_GREEN
            bar.line.width = P(1.5)
            num_color = LIME_GREEN
        else:
            bar.fill.fore_color.rgb = URGENCY_CARD
            bar.line.fill.background()
            num_color = URGENCY_SUB
        add_text(slide, str(n), sx, base - 0.30, step_w, 0.28, 10, num_color, bold=True,
                 align=PP_ALIGN.CENTER, fit=False)


def build_slide1(prs):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    add_header(slide, f"Where {FIRM_NAME} Stands Today", 9.30)

    quotes = [(who or "You said", q) for who, q in PROSPECT_QUOTES[:3] if q]

    # Left: "What you told us" — the prospect's own words, one card, one text box.
    if quotes:
        add_text(slide, "WHAT YOU TOLD US", 0.30, 1.08, 5.00, 0.20, 7.5, NAVY, bold=True,
                 fit=False)
        add_rect(slide, 0.30, 1.32, 5.00, 2.06, fill=WHITE)
        add_rect(slide, 0.30, 1.32, 0.05, 2.06, fill=LIME_GREEN)
        paras = []
        for i, (who, q) in enumerate(quotes):
            q = cap(q.strip().strip('"“”'), 130, f"PROSPECT_QUOTES[{i}]")
            paras.append(([(f"{who.upper()}:", 7.5, True, False, GREEN)], 0 if i == 0 else 9))
            paras.append(([(f"“{q}”", 11, False, True, NAVY)], 1))
        add_rich_text(slide, paras, 0.50, 1.40, 4.70, 1.92, label="PROSPECT_QUOTES",
                      anchor=MSO_ANCHOR.MIDDLE)

    # Right (or full width with no transcript): where they are vs. where they need to be.
    stair_x, stair_w = (5.50, 4.20) if quotes else (0.30, 9.40)
    add_text(slide, "WHERE YOU ARE TODAY", stair_x, 1.08, stair_w, 0.20, 7.5, NAVY,
             bold=True, fit=False)
    build_staircase_card(slide, stair_x, 1.32, stair_w, 2.06)

    # Bottom: the main gaps from the audit, grouped by category — one card,
    # three columns separated by thin rules, no per-item boxes.
    add_text(slide, "YOUR KEY GAPS", 0.30, 3.50, 4.00, 0.20, 7.5, NAVY, bold=True, fit=False)
    add_rect(slide, 0.30, 3.74, 9.40, 1.36, fill=WHITE)
    gaps = KEY_GAPS[:3]
    col_w = 9.40 / max(1, len(gaps))
    for i, (category, gap) in enumerate(gaps):
        cx = 0.30 + i * col_w
        if i:
            add_rule(slide, cx, 3.92, cx, 4.94, rgb("D9E2E8"))
        add_rect(slide, cx+0.24, 3.94, 0.28, 0.04, fill=LIME_GREEN if i % 2 else NAVY)
        add_text(slide, category.upper(), cx+0.18, 4.02, col_w-0.36, 0.22, 8, NAVY,
                 bold=True, max_lines=1, cap_label=f"KEY_GAPS[{i}] category")
        add_text(slide, gap, cx+0.18, 4.26, col_w-0.36, 0.76, 10.5, BODY_TEXT,
                 max_lines=3, cap_chars=90, cap_label=f"KEY_GAPS[{i}] gap")

    add_footer(slide, 1, 3, LOGO_PATH)


# ── Slide 2: Growth Plan ───────────────────────────────────────────

def build_slide2(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)

    # Background wash + header (shared with slides 1 and 3)
    add_header(slide, SLIDE_2_TITLE, 9.30)

    # Just the prospect's 3 priorities, full width. Each is one colored
    # header block over one white card holding a single bulleted text box.
    col_w, gap = 2.98, 0.23
    for col_i, (line1, line2, _hex, bullets) in enumerate(PRIORITIES[:3]):
        x = 0.30 + col_i * (col_w + gap)
        ac, title_color, kicker_color = PRIORITY_STYLE[col_i % len(PRIORITY_STYLE)]

        add_rect(slide, x, 1.10, col_w, 0.96, fill=ac)
        add_text(slide, f"PRIORITY {col_i+1}", x+0.20, 1.20, col_w-0.40, 0.20, 8,
                 kicker_color, bold=True, fit=False)
        # Title keeps the FILL's own line break (line1 / line2); both lines
        # share one size so the block reads as a single heading.
        title_w = col_w - 0.40
        sizes = [fit_text(t, title_w, 0.30, 14, bold=True, max_lines=1)[0] or 10.5
                 for t in (line1, line2) if t]
        title_size = min(sizes) if sizes else 14
        for row, t in enumerate([t for t in (line1, line2) if t]):
            add_text(slide, t, x+0.20, 1.40 + row*0.27, title_w, 0.30, title_size,
                     title_color, bold=True, max_lines=1, min_size=title_size,
                     cap_label=f"PRIORITIES[{col_i}] title")

        add_rect(slide, x, 2.06, col_w, 3.00, fill=WHITE)
        add_bullets(slide, [cap(b, 60, f"PRIORITIES[{col_i}] bullet {j+1}")
                            for j, b in enumerate(bullets[:5])],
                    x+0.14, 2.22, col_w-0.28, 2.74, 11, DARK_TEXT, ac,
                    label=f"PRIORITIES[{col_i}] bullets")

    add_footer(slide, 2, 3, LOGO_PATH)


# ── Slide 3: Investment & Next Steps ──────────────────────────────

def build_slide3(prs):
    # Never render an empty panel: with no transcript, Block 5 becomes
    # "What We Found" and this list still gets 3 research-based items.
    if len(WHAT_YOU_TOLD_US) != 3:
        raise SystemExit(
            f"WHAT_YOU_TOLD_US has {len(WHAT_YOU_TOLD_US)} items; Slide 3 needs exactly 3. "
            "With no discovery call transcript, copy the 3 research-based rows from "
            "section_11 Block 5 (What We Found).")

    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)

    # Background wash + header (mirrors slides 1 and 2)
    add_header(slide, "Your Investment & What Happens Next", 9.30)

    # Package cards — white with a thin colored left accent bar
    pkg_ys = [1.10, 2.25]
    for i, (label, price, retail, services, hex_c) in enumerate(PACKAGES[:2]):
        y = pkg_ys[i]
        ac = rgb(hex_c)
        add_rect(slide, 0.30, y, 4.42, 1.05, fill=WHITE)
        add_rect(slide, 0.30, y, 0.04, 1.05, fill=ac)
        add_text(slide, label, 0.56, y+0.12, 3.96, 0.20, 7.5, ac, bold=True,
                 max_lines=1, cap_label=f"PACKAGES[{i}] label")
        # Price, "/mo" and retail price share one baseline; "/mo" follows the
        # price's measured width so it never collides or wraps.
        price_top = y + 0.27
        price_size, _ = fit_text(price, 2.60, 0.58, 30, bold=True, max_lines=1)
        price_size = price_size or 22
        add_text(slide, price, 0.56, price_top, 2.60, 0.58, price_size, NAVY,
                 bold=True, fit=False)
        baseline = price_top + baseline_offset(price_size, bold=True)
        mo_x = 0.56 + _INSET_X + text_width(price, price_size, bold=True) + 0.04 - _INSET_X
        add_text(slide, "/mo", mo_x, baseline - baseline_offset(10), 0.50, 0.24, 10,
                 SLATE, wrap=False, fit=False)
        retail_x = mo_x + 2 * _INSET_X + text_width("/mo", 10) + 0.06
        add_text(slide, retail, retail_x, baseline - baseline_offset(10),
                 4.62 - retail_x, 0.24, 10, SLATE, align=PP_ALIGN.RIGHT,
                 max_lines=1, cap_label=f"PACKAGES[{i}] retail")
        add_text(slide, services, 0.56, y+0.80, 3.96, 0.22, 8.5, SLATE, max_lines=1,
                 cap_chars=61, cap_label=f"PACKAGES[{i}] services")

    # Bundle total — same 0.10in gap used between the two package cards above
    add_rect(slide, 0.30, 3.40, 4.42, 0.72, fill=NAVY)
    add_rect(slide, 0.30, 3.40, 0.04, 0.72, fill=LIME_GREEN)
    add_text(slide, "BUNDLE TOTAL", 0.56, 3.50, 1.70, 0.18, 7.5, LIME_GREEN, bold=True,
             fit=False)
    total_size, _ = fit_text(BUNDLE_TOTAL, 2.40, 0.40, 17, bold=True, max_lines=1)
    total_size = total_size or 13
    add_text(slide, BUNDLE_TOTAL, 0.56, 3.66, 2.40, 0.40, total_size, WHITE, bold=True,
             fit=False)
    # Savings line sits to the right of the measured total, on its baseline,
    # and shrinks to fit the space left inside the card.
    total_baseline = 3.66 + baseline_offset(total_size, bold=True)
    savings_x = 0.56 + text_width(BUNDLE_TOTAL, total_size, bold=True) + 0.20
    add_text(slide, BUNDLE_SAVINGS, savings_x, total_baseline - baseline_offset(8.5),
             4.62 - savings_x, 0.22, 8.5, BUNDLE_SAVINGS_TEXT, align=PP_ALIGN.RIGHT,
             max_lines=1, cap_label="BUNDLE_SAVINGS")

    # Ad spend note + ROI-on-ad-spend card — both omitted entirely when no
    # marketing/ads package was sold. Showing a return-on-ad-spend projection
    # for ad spend that isn't part of the recommendation is the same bug as
    # the missing-rationale note, just in a second spot on this slide.
    has_ad_spend = bool(AD_SPEND_NOTE)

    if has_ad_spend:
        add_rect(slide, 0.30, 4.22, 4.42, 0.42, fill=rgb("EEF2F5"))
        add_text(slide, AD_SPEND_NOTE, 0.50, 4.22, 4.02, 0.42, 8, BODY_TEXT,
                 cap_chars=120, cap_label="AD_SPEND_NOTE")

        # ROI card
        add_rect(slide, 5.00, 1.10, 4.70, 1.42, fill=ROI_BG)
        add_text(slide, "PROJECTED RETURN ON AD SPEND",
                 5.22, 1.26, 4.26, 0.16, 7.5, ROI_GREEN, bold=True)
        add_text(slide, "Average case value:", 5.22, 1.58, 1.95, 0.18, 8.5, DARK_TEXT, bold=True)
        add_text(slide, AVG_CASE_VALUE, 7.25, 1.58, 2.30, 0.18, 8.5, ROI_GREEN)
        add_rect(slide, 5.14, 1.83, 4.42, 0.28, fill=ROI_HILIGHT)
        add_text(slide, CONSERVATIVE_LABEL, 5.22, 1.90, 1.95, 0.18, 8.5, DARK_TEXT, bold=True)
        add_text(slide, CONSERVATIVE_RESULT, 7.25, 1.90, 2.30, 0.18, 8.5, ROI_GREEN, bold=True)
        add_text(slide, AGGRESSIVE_LABEL, 5.22, 2.21, 1.95, 0.18, 8.5, DARK_TEXT, bold=True)
        add_text(slide, AGGRESSIVE_RESULT, 7.25, 2.21, 2.30, 0.18, 8.5, ROI_GREEN, bold=True)

        told_us_y = 2.72
        timeline_ys = [2.99, 3.37, 3.75, 4.13]
    else:
        # No ROI card above it — give What You Told Us the full right-column height.
        told_us_y = 1.10
        timeline_ys = [1.46, 2.06, 2.66, 3.26]

    # How this plan addresses it — alternating light/white rows; the first dot pops in
    # lime green, the rest in ocean blue, so the lead item stands out
    add_text(slide, "HOW THIS PLAN ADDRESSES IT",
             5.00, told_us_y, 4.70, 0.18, 7.5, NAVY, bold=True)
    for i, (milestone, action) in enumerate(WHAT_YOU_TOLD_US[:3]):
        y = timeline_ys[i]
        row_bg = rgb("EEF2F5") if i % 2 == 0 else WHITE
        dot_color = LIME_GREEN if i == 0 else OCEAN_BLUE
        add_rect(slide, 5.00, y, 4.70, 0.29, fill=row_bg)
        add_rect(slide, 5.14, y+0.09, 0.10, 0.10, fill=dot_color)
        add_text(slide, milestone, 5.30, y, 1.25, 0.29, 8, NAVY, bold=True,
                 max_lines=1, anchor=MSO_ANCHOR.MIDDLE,
                 cap_chars=16, cap_label=f"WHAT_YOU_TOLD_US[{i}] label")
        add_text(slide, action, 6.50, y, 3.20, 0.29, 8, BODY_TEXT,
                 max_lines=1, anchor=MSO_ANCHOR.MIDDLE,
                 cap_chars=55, cap_label=f"WHAT_YOU_TOLD_US[{i}] action")

    # Closing quote bar — taller, darker navy, with the same lime accent stripe
    add_rect(slide, 0, 4.66, 10, 0.61, fill=rgb("002A41"))
    add_rect(slide, 0, 4.66, 10, 0.02, fill=LIME_GREEN)
    add_text(slide, CLOSING_QUOTE, 0.70, 4.70, 8.60, 0.54, 9.5, rgb("C2D8E6"),
             align=PP_ALIGN.CENTER, max_lines=2, anchor=MSO_ANCHOR.MIDDLE,
             cap_chars=200, cap_label="CLOSING_QUOTE")

    add_footer(slide, 3, 3, LOGO_PATH)


# ── Font embedding ────────────────────────────────────────────────
# python-pptx has no font-embedding API, so this patches the saved .pptx's
# raw OOXML directly. Mirrors what sales_companion_template.py already does
# for the PDF via reportlab's registerFont — without it, "Poppins" is just a
# name any viewer without the font installed (Google Slides, a PowerPoint
# missing the font) will silently substitute, which is why fonts can look
# inconsistent across machines even though the script never changes.

def embed_fonts(pptx_path, font_dir):
    import zipfile, shutil, tempfile

    faces = [
        ("regular", "Poppins-Regular.ttf"),
        ("bold",    "Poppins-Bold.ttf"),
        ("italic",  "Poppins-Italic.ttf"),
    ]
    faces = [(role, os.path.join(font_dir, fname)) for role, fname in faces
             if os.path.isfile(os.path.join(font_dir, fname))]
    if not faces:
        return

    NS = {
        "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
        "r":  "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
        "p":  "http://schemas.openxmlformats.org/presentationml/2006/main",
    }

    tmp_dir = tempfile.mkdtemp()
    try:
        with zipfile.ZipFile(pptx_path) as zin:
            zin.extractall(tmp_dir)

        # 1. Copy the font binaries into ppt/fonts/
        fonts_dir = os.path.join(tmp_dir, "ppt", "fonts")
        os.makedirs(fonts_dir, exist_ok=True)
        font_parts = []  # (role, part_filename)
        for i, (role, src_path) in enumerate(faces, start=1):
            part_name = f"font{i}.fntdata"
            shutil.copyfile(src_path, os.path.join(fonts_dir, part_name))
            font_parts.append((role, part_name))

        # 2. Register the .fntdata extension in [Content_Types].xml
        ct_path = os.path.join(tmp_dir, "[Content_Types].xml")
        ct_tree = etree.parse(ct_path)
        ct_root = ct_tree.getroot()
        if not ct_root.xpath("ct:Default[@Extension='fntdata']", namespaces=NS):
            default = etree.SubElement(ct_root, f"{{{NS['ct']}}}Default")
            default.set("Extension", "fntdata")
            default.set("ContentType", "application/x-fontdata")
        ct_tree.write(ct_path, xml_declaration=True, encoding="UTF-8", standalone=True)

        # 3. Add a relationship from presentation.xml to each font part
        rels_path = os.path.join(tmp_dir, "ppt", "_rels", "presentation.xml.rels")
        rels_tree = etree.parse(rels_path)
        rels_root = rels_tree.getroot()
        existing_ids = {el.get("Id") for el in rels_root}

        def new_rid(n=[1]):
            while f"rIdEmbedFont{n[0]}" in existing_ids:
                n[0] += 1
            rid = f"rIdEmbedFont{n[0]}"
            n[0] += 1
            return rid

        rel_ids = {}
        for role, part_name in font_parts:
            rid = new_rid()
            rel = etree.SubElement(rels_root, f"{{{NS['pr']}}}Relationship")
            rel.set("Id", rid)
            rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font")
            rel.set("Target", f"fonts/{part_name}")
            rel_ids[role] = rid
        rels_tree.write(rels_path, xml_declaration=True, encoding="UTF-8", standalone=True)

        # 4. Declare the embedded font in presentation.xml and flip the
        #    embedTrueTypeFonts switch on.
        pres_path = os.path.join(tmp_dir, "ppt", "presentation.xml")
        pres_tree = etree.parse(pres_path)
        pres_root = pres_tree.getroot()
        pres_root.set("embedTrueTypeFonts", "1")

        embedded_font = etree.Element(f"{{{NS['p']}}}embeddedFont")
        font_el = etree.SubElement(embedded_font, f"{{{NS['p']}}}font")
        font_el.set("typeface", FONT)
        for role, rid in rel_ids.items():
            face_el = etree.SubElement(embedded_font, f"{{{NS['p']}}}{role}")
            face_el.set(f"{{{NS['r']}}}id", rid)

        embedded_font_lst = etree.Element(f"{{{NS['p']}}}embeddedFontLst")
        embedded_font_lst.append(embedded_font)

        # Schema order (CT_Presentation) requires embeddedFontLst to sit
        # after notesSz and before defaultTextStyle/custShowLst/etc.
        notes_sz = pres_root.find(f"{{{NS['p']}}}notesSz")
        default_text_style = pres_root.find(f"{{{NS['p']}}}defaultTextStyle")
        if notes_sz is not None:
            insert_at = list(pres_root).index(notes_sz) + 1
        elif default_text_style is not None:
            insert_at = list(pres_root).index(default_text_style)
        else:
            insert_at = len(pres_root)
        pres_root.insert(insert_at, embedded_font_lst)

        pres_tree.write(pres_path, xml_declaration=True, encoding="UTF-8", standalone=True)

        # 5. Re-zip in place
        tmp_pptx = pptx_path + ".tmp"
        with zipfile.ZipFile(tmp_pptx, "w", zipfile.ZIP_DEFLATED) as zout:
            for root_dir, _, files in os.walk(tmp_dir):
                for fname in files:
                    full = os.path.join(root_dir, fname)
                    arcname = os.path.relpath(full, tmp_dir)
                    zout.write(full, arcname)
        shutil.move(tmp_pptx, pptx_path)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


# ── Assemble ──────────────────────────────────────────────────────

prs = Presentation()
prs.slide_width  = Inches(10)
prs.slide_height = Inches(5.625)

build_slide1(prs)
build_slide2(prs)
build_slide3(prs)

os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True) if os.path.dirname(OUTPUT_PATH) else None
prs.save(OUTPUT_PATH)
embed_fonts(OUTPUT_PATH, FONT_DIR)

if _SHRUNK:
    print(f"NOTE: {len(_SHRUNK)} field(s) were shrunk to fit their box:")
    for label, before, after in _SHRUNK:
        print(f"  - {label}: {before:g}pt → {after:g}pt")

if _TRUNCATED:
    print(f"WARNING: {len(_TRUNCATED)} field(s) exceeded their character budget "
          f"and were truncated — consider shortening the source copy instead:")
    for label, original in _TRUNCATED:
        print(f"  - {label}: {original!r}")

print(f"Saved: {OUTPUT_PATH}  ({len(prs.slides)} slides)")
