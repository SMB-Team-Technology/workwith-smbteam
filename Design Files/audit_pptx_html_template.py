"""
Audit PowerPoint Template — SMB Team (HTML/CSS rendering engine)
=================================================================
Generates a 3-slide proposal PPTX from the Growth, Profit, and Freedom Roadmap.

Unlike audit_pptx_template.py (which places python-pptx text boxes and
rectangles by hand), this version builds each slide as an HTML/CSS document —
real flexbox/grid layout, gradients, shadows, and CSS text truncation — then
renders it to a high-resolution PNG with headless Chromium (Playwright) and
drops that image full-bleed into a PPTX slide. The .pptx output, filenames,
and downstream pipeline (email attachments, Slack, HubSpot) are unchanged;
only how the slide content gets drawn changes.

Trade-off: slides are rasterized images, not editable PowerPoint text boxes.
Fonts are baked into the render (no runtime font-embedding step needed), and
CSS handles overflow/line-clamping instead of the character-budget guesswork
the shape-based engine needed.

Do not modify the RENDER ENGINE below the FILL section.
Only replace the # FILL: placeholders with audit-specific content.

Slide 1 — Where [Firm] Stands Today         (assessment overview)
Slide 2 — Your Growth Plan: 3 Priorities    (action plan)
Slide 3 — Your Investment & What's Next     (pricing + first 90 days)

Output: [friendly-name]/[FirmName]_[Date]_Proposal.pptx
"""

import base64
import html
import os

from pptx import Presentation
from pptx.util import Inches

# ═══════════════════════════════════════════════════════════════════
# FILL — Replace every placeholder. Do not delete any variable.
# (Same shape as audit_pptx_template.py — copy values across unchanged.)
# ═══════════════════════════════════════════════════════════════════

# Report metadata
FIRM_NAME      = "FIRM NAME HERE"                    # FILL: full firm name
SALES_REP      = "Sales Rep Name"                    # FILL: rep's full name
OUTPUT_PATH    = "friendly-name/FirmName_Date_Proposal.pptx"  # FILL: output path

# Optional images
WEBSITE_SCREENSHOT_PATH = None   # Set to a local PNG path to show a screenshot; None = skip entirely

# ── Slide 1 ──────────────────────────────────────────────────────
URGENCY_SCORE = "7"                          # FILL: copy from section_05 urgency score

# Pillar status — each is ("RED" | "AMBER" | "GREEN", label, one-line detail)
PILLARS = [
    ("RED",   "CRITICAL", "100% from referrals"),   # FILL: Lead Generation
    ("AMBER", "AMBER",    "Leads falling through"),  # FILL: Intake
    ("AMBER", "AMBER",    "No KPIs / bottleneck"),   # FILL: Team
    ("AMBER", "AMBER",    "No forecasting"),          # FILL: Profit Plan
]
PILLAR_NAMES = ["Lead Generation", "Intake", "Team", "Profit Plan"]

# Key findings — list of ("neg"|"pos", "one-sentence finding")
# Use "neg" for problems (red ✕), "pos" for strengths (green ✓).
# Include 3–4 total; at least one "pos" if a genuine strength exists.
FINDINGS = [
    ("neg", "Finding one — consequence-first, specific to this firm"),
    ("neg", "Finding two — consequence-first, specific to this firm"),
    ("neg", "Finding three — consequence-first, specific to this firm"),
    ("pos", "Finding four — genuine strength of this firm"),
]

# Competitor table — list of (name, "XXX reviews", "brief note")
# List 3 competitors + the client row is added automatically.
COMPETITORS = [
    ("Competitor One",   "XXX reviews", "detail · note"),
    ("Competitor Two",   "XXX reviews", "detail · note"),
    ("Competitor Three", "XXX reviews", "detail · note"),
]
CLIENT_REVIEWS      = "0 reviews"      # FILL: client's actual review count
CLIENT_REVIEWS_NOTE = "← You are here"

# Stage strip (right panel, Slide 1)
STAGE_TEXT = "Stage 4: Small Business Manager  →  Goal: Stage 6, Law Firm Owner"  # FILL

# ── Slide 2 ──────────────────────────────────────────────────────
SLIDE_2_TITLE = "Your Growth Plan: 3 Priorities to Reach $X,XXX/Month"  # FILL

SMB_MODEL_DESC = (
    "All four pillars must work together. "
    "Missing any one means growth stalls regardless of ad spend."
)  # FILL: brief model description

GOAL_HEADLINE = "$X → $Y revenue"       # FILL: e.g. "$1M → $1.8M revenue"
GOAL_DBM      = "Owner takes real time off"  # FILL: DBM outcome phrase

# Each priority: (line1, line2, accent_color_hex, [5 bullet strings])
# Titles must reflect what this firm is actually buying and its real priorities —
# see .claude/commands/audit-pptx.md. Do not default a column to "Marketing Engine"
# unless a marketing/ads package is actually part of this firm's recommendation.
PRIORITIES = [
    (
        "Priority", "One", "0091C9",
        [
            "Bullet 1 — specific action for this firm",
            "Bullet 2 — specific action for this firm",
            "Bullet 3 — specific action for this firm",
            "Bullet 4 — specific action for this firm",
            "Bullet 5 — specific action for this firm",
        ],
    ),
    (
        "Fix Intake &", "Stop Losing Cases", "69CD2B",
        [
            "Bullet 1 — specific action for this firm",
            "Bullet 2 — specific action for this firm",
            "Bullet 3 — specific action for this firm",
            "Bullet 4 — specific action for this firm",
            "Bullet 5 — specific action for this firm",
        ],
    ),
    (
        "Install Team &", "Profit Systems", "003A59",
        [
            "Bullet 1 — specific action for this firm",
            "Bullet 2 — specific action for this firm",
            "Bullet 3 — specific action for this firm",
            "Bullet 4 — specific action for this firm",
            "Bullet 5 — specific action for this firm",
        ],
    ),
]

# ── Slide 3 ──────────────────────────────────────────────────────
# Package cards — (label, bundled_price, retail_price, services_line, accent_color_hex)
PACKAGES = [
    (
        "FULL SERVICE MARKETING — GROWTH",
        "$X,XXX", "$X,XXX/mo",
        "Website · Google Ads · LSA · Meta Ads · GBP optimization",
        "0091C9",
    ),
    (
        "ELITE COACH PLUS",
        "$X,XXX", "$X,XXX/mo",
        "Weekly group coaching · KPI scorecards · Intake framework",
        "003A59",
    ),
]

BUNDLE_TOTAL   = "$X,XXX / mo"                      # FILL: sum of bundled prices
BUNDLE_SAVINGS = "Save $X,XXX/mo by bundling"        # FILL

AD_SPEND_NOTE = (
    "+ Recommended ad spend: $X,XXX–$XX,XXX/mo paid directly to Google/Meta"
)  # FILL — set to None if no marketing/ads package is part of this firm's recommendation

AVG_CASE_VALUE     = "$X,XXX"                         # FILL
CONSERVATIVE_LABEL = "Conservative  (X cases/mo):"   # FILL
CONSERVATIVE_RESULT = "$XX,XXX revenue · X.X× ROAS"  # FILL
AGGRESSIVE_LABEL   = "Aggressive  (X cases/mo):"     # FILL
AGGRESSIVE_RESULT  = "$XX,XXX revenue · X.X× ROAS"   # FILL

# Timeline — 5 items: (milestone_label, action_text)
TIMELINE = [
    ("Day 1",   "Action at day 1 — specific to this firm"),
    ("Day 14",  "Action at day 14 — specific to this firm"),
    ("Week 2",  "Action at week 2 — specific to this firm"),
    ("Week 3",  "Action at week 3 — specific to this firm"),
    ("Month 3", "Action at month 3 — specific to this firm"),
]

CLOSING_QUOTE = (
    '"Closing quote tied to this firm\'s DBM — exact transcript words '
    'or a sharp synthesis of the central opportunity."'
)  # FILL

# ═══════════════════════════════════════════════════════════════════
# RENDER ENGINE — DO NOT MODIFY BELOW THIS LINE
# ═══════════════════════════════════════════════════════════════════

SLIDE_W_PX, SLIDE_H_PX = 1920, 1080  # 16:9 at 192 DPI for a 10in x 5.625in slide

_HERE = os.path.dirname(__file__)
# This script is copied into each firm's own folder (a sibling of "Design
# Files/" at the repo root) to run — assets are addressed relative to the
# repo root instead, since only this .py file (and the logo, copied
# alongside it by the audit-pptx skill) travel with the firm's copy.
_DESIGN_DIR = os.path.join(_HERE, "..", "Design Files")
FONT_DIR = os.path.join(_DESIGN_DIR, "fonts")
LOGO_PATH = os.path.join(_HERE, "smb_team_logo.png")


def _b64(path):
    with open(path, "rb") as f:
        return base64.b64encode(f.read()).decode("ascii")


def _data_uri(path, mime):
    return f"data:{mime};base64,{_b64(path)}"


def _font_face_css():
    faces = [
        ("Poppins", 400, "normal", "Poppins-Regular.ttf"),
        ("Poppins", 700, "normal", "Poppins-Bold.ttf"),
        ("Poppins", 400, "italic", "Poppins-Italic.ttf"),
    ]
    rules = []
    for family, weight, style, fname in faces:
        path = os.path.join(FONT_DIR, fname)
        if not os.path.isfile(path):
            continue
        uri = _data_uri(path, "font/ttf")
        rules.append(
            f"@font-face {{ font-family:'{family}'; font-weight:{weight}; "
            f"font-style:{style}; src:url('{uri}') format('truetype'); }}"
        )
    return "\n".join(rules)


def _logo_uri():
    return _data_uri(LOGO_PATH, "image/png") if os.path.isfile(LOGO_PATH) else ""


def esc(text):
    return html.escape("" if text is None else str(text))


# ── Shared CSS ────────────────────────────────────────────────────

BASE_CSS = """
:root {
  --navy:#003A59; --dark-navy:#265872; --ocean:#0091C9; --lime:#69CD2B;
  --red:#C0392B; --amber:#D97706; --green:#16A34A; --slate:#64748B;
  --light-blue:#8BADD0; --footer-text:#5A7A9A; --neg-bg:#2C1010;
  --pos-bg:#0C2818; --neg-text:#F0BABA; --pos-text:#86EFAC;
  --comp-alt:#F0F4FA; --roi-bg:#ECFDF5; --roi-green:#065F46;
  --roi-hi:#D1FAE5; --bundle-sub:#7CA0C0; --strike:#334155;
  --ink:#1E293B; --white:#fff;
}
* { box-sizing:border-box; margin:0; padding:0; }
html, body { width:1920px; height:1080px; overflow:hidden; }
body { font-family:'Poppins', sans-serif; color:var(--ink); background:var(--white); }
.slide { position:relative; width:1920px; height:1080px; }

.banner {
  position:absolute; top:0; left:0; right:0; height:170px;
  background:linear-gradient(120deg, var(--navy) 0%, var(--dark-navy) 100%);
  display:flex; flex-direction:column; justify-content:center; padding:0 56px;
}
.banner .kicker {
  color:var(--lime); font-weight:700; font-size:22px; letter-spacing:2px;
  text-transform:uppercase; margin-bottom:10px;
}
.banner .title { color:var(--white); font-weight:700; font-size:52px; line-height:1.15; }

.footer {
  position:absolute; left:0; right:0; bottom:0; height:64px;
  background:var(--navy); display:flex; align-items:center; padding:0 42px;
}
.footer img { height:38px; margin-right:24px; }
.footer .meta { color:var(--footer-text); font-size:16px; flex:1; }
.footer .page { color:var(--footer-text); font-size:16px; }

.clamp1 { display:-webkit-box; -webkit-line-clamp:1; -webkit-box-orient:vertical; overflow:hidden; }
.clamp2 { display:-webkit-box; -webkit-line-clamp:2; -webkit-box-orient:vertical; overflow:hidden; }
.ellipsis { white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }
"""


def _base_html(body_html, extra_css=""):
    return f"""<!doctype html>
<html><head><meta charset="utf-8">
<style>
{_font_face_css()}
{BASE_CSS}
{extra_css}
</style></head>
<body>{body_html}</body></html>"""


def _footer_html(page, total):
    logo_uri = _logo_uri()
    logo_tag = f'<img src="{logo_uri}">' if logo_uri else ""
    return f"""
    <div class="footer">
      {logo_tag}
      <div class="meta">CONFIDENTIAL &nbsp;·&nbsp; Prepared by {esc(SALES_REP)} | SMB Team</div>
      <div class="page">{page} / {total}</div>
    </div>"""


STATUS_COLOR_VAR = {"RED": "var(--red)", "AMBER": "var(--amber)", "GREEN": "var(--green)"}
PRIORITY_LIGHT = {
    "0091C9": "#E3F4FA",
    "69CD2B": "#EDF9E6",
    "003A59": "#E0E7EB",
}
ACCENT_TEXT_COLOR = {"69CD2B": "var(--navy)"}


# ── Slide 1: Assessment Overview ──────────────────────────────────

def build_slide1_html():
    pillar_cards = ""
    for i, (status, label, detail) in enumerate(PILLARS):
        sc = STATUS_COLOR_VAR[status]
        pillar_cards += f"""
        <div class="pillar">
          <div class="pillar-bar" style="background:{sc};"></div>
          <div class="pillar-body">
            <div class="pillar-name">{esc(PILLAR_NAMES[i])}</div>
            <div class="pillar-label" style="color:{sc};">{esc(label)}</div>
            <div class="pillar-detail clamp1">{esc(detail)}</div>
          </div>
        </div>"""

    finding_rows = ""
    for ftype, text in FINDINGS[:4]:
        is_neg = ftype == "neg"
        bg   = "var(--neg-bg)" if is_neg else "var(--pos-bg)"
        dot  = "var(--red)" if is_neg else "var(--green)"
        sym  = "&#10005;" if is_neg else "&#10003;"
        txt_color = "var(--neg-text)" if is_neg else "var(--pos-text)"
        finding_rows += f"""
        <div class="finding" style="background:{bg};">
          <div class="finding-dot" style="background:{dot};">{sym}</div>
          <div class="finding-text clamp2" style="color:{txt_color};">{esc(text)}</div>
        </div>"""

    screenshot_html = ""
    if WEBSITE_SCREENSHOT_PATH and os.path.isfile(WEBSITE_SCREENSHOT_PATH):
        uri = _data_uri(WEBSITE_SCREENSHOT_PATH, "image/png")
        screenshot_html = f'<img class="site-shot" src="{uri}">'

    comp_rows = ""
    for i, (name, reviews, detail) in enumerate(COMPETITORS[:3]):
        bg = "var(--comp-alt)" if i % 2 == 0 else "var(--white)"
        comp_rows += f"""
        <div class="comp-row" style="background:{bg};">
          <div class="comp-name ellipsis">{esc(name)}</div>
          <div class="comp-reviews ellipsis">{esc(reviews)}</div>
          <div class="comp-detail ellipsis">{esc(detail)}</div>
        </div>"""

    extra_css = """
    .body-area { position:absolute; top:170px; left:0; right:0; bottom:64px; display:flex; }
    .left-col { width:58%; padding:28px 20px 28px 56px; display:flex; flex-direction:column; }
    .right-col { width:42%; background:var(--white); padding:28px 44px 28px 32px;
                 display:flex; flex-direction:column; box-shadow:-8px 0 24px rgba(0,0,0,0.06); }

    .top-row { display:flex; gap:18px; align-items:stretch; margin-bottom:22px; }
    .urgency-box {
      width:150px; border-radius:14px; background:linear-gradient(160deg,#D9483A,var(--red));
      color:var(--white); display:flex; flex-direction:column; align-items:center;
      justify-content:center; padding:10px; box-shadow:0 6px 18px rgba(192,57,43,0.35);
    }
    .urgency-box .num { font-size:56px; font-weight:700; line-height:1; }
    .urgency-box .lbl { font-size:13px; font-weight:700; text-align:center; margin-top:6px; letter-spacing:0.5px; }

    .pillars { flex:1; display:grid; grid-template-columns:repeat(4,1fr); gap:12px; }
    .pillar { border-radius:12px; overflow:hidden; background:var(--dark-navy);
              box-shadow:0 4px 10px rgba(0,0,0,0.15); display:flex; flex-direction:column; }
    .pillar-bar { height:8px; }
    .pillar-body { padding:12px 12px; color:var(--white); }
    .pillar-name { font-size:15px; font-weight:700; margin-bottom:6px; }
    .pillar-label { font-size:13px; font-weight:700; margin-bottom:6px; }
    .pillar-detail { font-size:13px; color:var(--light-blue); }

    .findings-label { color:var(--lime); font-weight:700; font-size:16px;
                       letter-spacing:1px; margin-bottom:12px; }
    .findings { display:flex; flex-direction:column; gap:14px; flex:1; }
    .finding { flex:1; border-radius:10px; padding:14px 18px; display:flex; align-items:center; gap:14px; }
    .finding-dot { min-width:34px; height:34px; border-radius:50%; color:var(--white);
                   display:flex; align-items:center; justify-content:center; font-size:16px; font-weight:700; }
    .finding-text { font-size:18px; line-height:1.3; }

    .here-strip { border-radius:12px; background:linear-gradient(120deg,var(--navy),var(--dark-navy));
                  padding:16px 20px; display:flex; flex-direction:column; gap:6px;
                  border-left:6px solid var(--ocean); margin-bottom:18px; }
    .here-strip .kicker2 { color:var(--lime); font-weight:700; font-size:14px; letter-spacing:1px; }
    .here-strip .stage { color:var(--white); font-weight:700; font-size:19px; }

    .site-shot { width:100%; border-radius:10px; margin-bottom:18px; box-shadow:0 4px 14px rgba(0,0,0,0.15); }

    .comp-label { color:var(--slate); font-weight:700; font-size:14px; letter-spacing:1px; margin-bottom:10px; }
    .comp-table { border-radius:10px; overflow:hidden; border:1px solid #E5EAF1; }
    .comp-row { display:flex; align-items:center; padding:12px 16px; gap:12px; }
    .comp-name { flex:2; font-size:15px; font-weight:600; }
    .comp-reviews { flex:1; font-size:14px; color:var(--green); font-weight:600; }
    .comp-detail { flex:1.4; font-size:13px; color:var(--slate); text-align:right; }

    .client-row { display:flex; align-items:center; padding:14px 16px; gap:12px; margin-top:8px;
                  border-radius:10px; background:#E3F4FA; border-left:6px solid var(--ocean); }
    .client-row .comp-name { color:var(--navy); font-weight:700; }
    .client-row .comp-reviews { color:var(--ocean); font-weight:700; }
    """

    body = f"""
    <div class="slide">
      <div class="banner">
        <div class="kicker">Law Firm Growth Audit &nbsp;·&nbsp; {esc(FIRM_NAME.upper())}</div>
        <div class="title">Where {esc(FIRM_NAME)} Stands Today</div>
      </div>
      <div class="body-area">
        <div class="left-col">
          <div class="top-row">
            <div class="urgency-box">
              <div class="num">{esc(URGENCY_SCORE)}</div>
              <div class="lbl">COMPETITIVE<br>URGENCY</div>
            </div>
            <div class="pillars">{pillar_cards}</div>
          </div>
          <div class="findings-label">KEY FINDINGS</div>
          <div class="findings">{finding_rows}</div>
        </div>
        <div class="right-col">
          <div class="here-strip">
            <div class="kicker2">YOU ARE HERE</div>
            <div class="stage">{esc(STAGE_TEXT)}</div>
          </div>
          {screenshot_html}
          <div class="comp-label">COMPETITOR LANDSCAPE</div>
          <div class="comp-table">{comp_rows}</div>
          <div class="client-row">
            <div class="comp-name ellipsis" style="flex:2;">{esc(FIRM_NAME)}</div>
            <div class="comp-reviews ellipsis" style="flex:1;">{esc(CLIENT_REVIEWS)}</div>
            <div class="comp-detail ellipsis" style="flex:1.4;">{esc(CLIENT_REVIEWS_NOTE)}</div>
          </div>
        </div>
      </div>
      {_footer_html(1, 3)}
    </div>"""
    return _base_html(body, extra_css)


# ── Slide 2: Growth Plan ───────────────────────────────────────────

def build_slide2_html():
    priority_cols = ""
    for col_i, (line1, line2, hex_color, bullets) in enumerate(PRIORITIES):
        ac = f"#{hex_color}"
        light = PRIORITY_LIGHT.get(hex_color, "#F8F8FF")
        header_text = ACCENT_TEXT_COLOR.get(hex_color, "var(--white)")
        bullet_rows = ""
        for row_i, bullet in enumerate(bullets[:5]):
            bg = light if row_i % 2 == 0 else "var(--white)"
            bullet_rows += f"""
            <div class="bullet-row" style="background:{bg};">
              <div class="bullet-dot" style="background:{ac};"></div>
              <div class="bullet-text clamp2">{esc(bullet)}</div>
            </div>"""
        priority_cols += f"""
        <div class="priority-col">
          <div class="priority-header" style="background:{ac}; color:{header_text};">
            <div class="priority-num">0{col_i + 1}</div>
            <div class="priority-title">{esc(line1)}<br>{esc(line2)}</div>
          </div>
          <div class="bullet-list">{bullet_rows}</div>
        </div>"""

    extra_css = """
    .body-area { position:absolute; top:170px; left:0; right:0; bottom:64px; display:flex; gap:20px; padding:26px 56px; }
    .model-panel { width:340px; border-radius:14px; background:linear-gradient(160deg,var(--navy),var(--dark-navy));
                   padding:26px 24px; display:flex; flex-direction:column; box-shadow:0 6px 18px rgba(0,0,0,0.18); }
    .model-panel .kicker { color:var(--lime); font-weight:700; font-size:15px; letter-spacing:1px; margin-bottom:14px; }
    .model-panel .desc { color:#A8BFDA; font-size:16px; line-height:1.5; flex:1; }
    .goal-card { border-radius:12px; background:var(--lime); padding:18px 20px; margin-top:16px; }
    .goal-card .headline { color:var(--navy); font-weight:700; font-size:22px; margin-bottom:6px; }
    .goal-card .dbm { color:var(--navy); font-size:15px; }

    .priorities { flex:1; display:flex; gap:20px; }
    .priority-col { flex:1; display:flex; flex-direction:column; border-radius:14px; overflow:hidden;
                    box-shadow:0 4px 14px rgba(0,0,0,0.12); }
    .priority-header { padding:18px 20px; }
    .priority-num { font-size:20px; font-weight:700; opacity:0.85; margin-bottom:6px; }
    .priority-title { font-size:22px; font-weight:700; line-height:1.3; }
    .bullet-list { flex:1; display:flex; flex-direction:column; }
    .bullet-row { flex:1; display:flex; align-items:center; gap:14px; padding:12px 18px; }
    .bullet-dot { width:12px; height:12px; border-radius:3px; flex-shrink:0; }
    .bullet-text { font-size:15px; line-height:1.35; color:var(--ink); }
    """

    body = f"""
    <div class="slide">
      <div class="banner">
        <div class="kicker">Law Firm Growth Audit &nbsp;·&nbsp; {esc(FIRM_NAME.upper())}</div>
        <div class="title">{esc(SLIDE_2_TITLE)}</div>
      </div>
      <div class="body-area">
        <div class="model-panel">
          <div class="kicker">THE SMB TEAM MODEL</div>
          <div class="desc">{esc(SMB_MODEL_DESC)}</div>
          <div class="goal-card">
            <div class="headline">{esc(GOAL_HEADLINE)}</div>
            <div class="dbm">{esc(GOAL_DBM)}</div>
          </div>
        </div>
        <div class="priorities">{priority_cols}</div>
      </div>
      {_footer_html(2, 3)}
    </div>"""
    return _base_html(body, extra_css)


# ── Slide 3: Investment & Next Steps ──────────────────────────────

def build_slide3_html():
    pkg_cards = ""
    for label, price, retail, services, hex_c in PACKAGES[:2]:
        ac = f"#{hex_c}"
        pkg_cards += f"""
        <div class="pkg-card" style="border-left-color:{ac};">
          <div class="pkg-label" style="color:{ac};">{esc(label)}</div>
          <div class="pkg-price-row">
            <span class="pkg-price">{esc(price)}</span><span class="pkg-permo">/mo</span>
            <span class="pkg-retail">{esc(retail)}</span>
          </div>
          <div class="pkg-services ellipsis">{esc(services)}</div>
        </div>"""

    # Omitted entirely when no marketing/ads package was sold
    ad_spend_html = f'<div class="ad-spend">{esc(AD_SPEND_NOTE)}</div>' if AD_SPEND_NOTE else ""

    timeline_rows = ""
    for milestone, action in TIMELINE[:5]:
        timeline_rows += f"""
        <div class="tl-row">
          <div class="tl-dot"></div>
          <div class="tl-milestone">{esc(milestone)}</div>
          <div class="tl-action clamp1">{esc(action)}</div>
        </div>"""

    extra_css = """
    .body-area { position:absolute; top:170px; left:0; right:0; bottom:110px; display:flex; gap:24px; padding:26px 56px 0; }
    .left-col { width:47%; display:flex; flex-direction:column; gap:16px; }
    .right-col { flex:1; display:flex; flex-direction:column; gap:24px; align-content:flex-start; }

    .pkg-card { background:var(--white); border:1px solid #E5EAF1; border-left-width:8px; border-left-style:solid;
                border-radius:12px; padding:18px 22px; box-shadow:0 3px 10px rgba(0,0,0,0.06); }
    .pkg-label { font-weight:700; font-size:14px; letter-spacing:0.5px; margin-bottom:8px; }
    .pkg-price-row { display:flex; align-items:baseline; gap:10px; margin-bottom:8px; }
    .pkg-price { font-size:40px; font-weight:700; color:var(--navy); }
    .pkg-permo { font-size:16px; color:var(--slate); }
    .pkg-retail { font-size:16px; color:var(--strike); text-decoration:line-through; margin-left:auto; }
    .pkg-services { font-size:14px; color:var(--slate); }

    .bundle-total { border-radius:12px; background:linear-gradient(120deg,var(--navy),var(--dark-navy));
                    padding:18px 22px; display:flex; align-items:center; gap:20px; }
    .bundle-total .lbl { color:var(--bundle-sub); font-weight:700; font-size:14px; }
    .bundle-total .val { color:var(--ocean); font-weight:700; font-size:30px; }
    .bundle-total .savings { color:var(--bundle-sub); font-size:14px; margin-left:auto; text-align:right; }

    .ad-spend { border-radius:10px; background:#EEF2F8; padding:14px 20px; font-size:14px; color:var(--slate); }

    .roi-card { border-radius:12px; background:var(--roi-bg); padding:20px 24px; }
    .roi-card .title { color:var(--roi-green); font-weight:700; font-size:14px; letter-spacing:0.5px; margin-bottom:12px; }
    .roi-line { display:flex; justify-content:space-between; font-size:16px; padding:6px 0; }
    .roi-hi { background:var(--roi-hi); border-radius:8px; padding:8px 14px; margin:4px 0; }
    .roi-hi .roi-line { padding:0; }
    .roi-strong { font-weight:700; color:var(--roi-green); }

    .tl-title { color:var(--navy); font-weight:700; font-size:14px; letter-spacing:0.5px; margin-bottom:6px; }
    .timeline { display:flex; flex-direction:column; gap:26px; }
    .tl-row { display:flex; align-items:center; gap:16px; }
    .tl-dot { width:14px; height:14px; border-radius:50%; background:var(--navy); flex-shrink:0; }
    .tl-milestone { width:90px; font-weight:700; color:var(--navy); font-size:15px; }
    .tl-action { flex:1; font-size:15px; color:var(--ink); }

    .closing-bar { position:absolute; left:0; right:0; bottom:64px; height:46px; background:var(--navy);
                   display:flex; align-items:center; justify-content:center; padding:0 60px; }
    .closing-bar .quote { color:var(--light-blue); font-style:italic; font-size:17px; text-align:center; }
    """

    body = f"""
    <div class="slide">
      <div class="banner">
        <div class="kicker">Law Firm Growth Audit &nbsp;·&nbsp; {esc(FIRM_NAME.upper())}</div>
        <div class="title">Your Investment &amp; What Happens Next</div>
      </div>
      <div class="body-area">
        <div class="left-col">
          {pkg_cards}
          <div class="bundle-total">
            <div>
              <div class="lbl">BUNDLE TOTAL</div>
              <div class="val">{esc(BUNDLE_TOTAL)}</div>
            </div>
            <div class="savings">{esc(BUNDLE_SAVINGS)}</div>
          </div>
          {ad_spend_html}
        </div>
        <div class="right-col">
          <div class="roi-card">
            <div class="title">PROJECTED RETURN ON AD SPEND</div>
            <div class="roi-line"><span>Average case value:</span><span class="roi-strong">{esc(AVG_CASE_VALUE)}</span></div>
            <div class="roi-hi"><div class="roi-line"><span>{esc(CONSERVATIVE_LABEL)}</span><span class="roi-strong">{esc(CONSERVATIVE_RESULT)}</span></div></div>
            <div class="roi-line"><span>{esc(AGGRESSIVE_LABEL)}</span><span class="roi-strong">{esc(AGGRESSIVE_RESULT)}</span></div>
          </div>
          <div class="tl-title">WHAT HAPPENS IN THE FIRST 90 DAYS</div>
          <div class="timeline">{timeline_rows}</div>
        </div>
      </div>
      <div class="closing-bar"><div class="quote">{esc(CLOSING_QUOTE)}</div></div>
      {_footer_html(3, 3)}
    </div>"""
    return _base_html(body, extra_css)


# ── Chromium rendering ──────────────────────────────────────────────

def _launch_browser(p):
    # In CI (after `playwright install chromium`) the default launch works.
    # PLAYWRIGHT_CHROMIUM_PATH lets a preinstalled-but-differently-versioned
    # Chromium (e.g. this sandbox's /opt/pw-browsers) be used instead of
    # re-downloading a browser build.
    override = os.environ.get("PLAYWRIGHT_CHROMIUM_PATH")
    if override:
        return p.chromium.launch(executable_path=override)
    return p.chromium.launch()


def render_slides_to_png(html_docs, out_dir):
    from playwright.sync_api import sync_playwright

    paths = []
    with sync_playwright() as p:
        browser = _launch_browser(p)
        page = browser.new_page(viewport={"width": SLIDE_W_PX, "height": SLIDE_H_PX})
        for i, doc in enumerate(html_docs, start=1):
            page.set_content(doc, wait_until="load")
            out_path = os.path.join(out_dir, f"_slide{i}.png")
            page.screenshot(path=out_path)
            paths.append(out_path)
        browser.close()
    return paths


# ── Assemble ──────────────────────────────────────────────────────

def main():
    out_dir = os.path.dirname(OUTPUT_PATH) or "."
    os.makedirs(out_dir, exist_ok=True)

    html_docs = [build_slide1_html(), build_slide2_html(), build_slide3_html()]
    png_paths = render_slides_to_png(html_docs, out_dir)

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    blank_layout = prs.slide_layouts[6]

    for png_path in png_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(png_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    prs.save(OUTPUT_PATH)

    for png_path in png_paths:
        os.remove(png_path)

    print(f"Saved: {OUTPUT_PATH}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
