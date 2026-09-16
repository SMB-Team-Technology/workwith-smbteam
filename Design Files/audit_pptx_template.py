"""
Audit PowerPoint Template — SMB Team
======================================
Generates a 3-slide proposal PPTX from the Growth, Profit, and Freedom Roadmap.
Uses python-pptx. Do not modify the layout engine below the FILL section.
Only replace the # FILL: placeholders with audit-specific content.

Slide 1 — Where [Firm] Stands Today         (assessment overview)
Slide 2 — Your Growth Plan: 3 Priorities    (action plan)
Slide 3 — Your Investment & What's Next     (pricing + first 90 days)

Output: [friendly-name]/[FirmName]_[Date]_Proposal.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from lxml import etree
import os

# ═══════════════════════════════════════════════════════════════════
# FILL — Replace every placeholder. Do not delete any variable.
# ═══════════════════════════════════════════════════════════════════

# Report metadata
FIRM_NAME      = "FIRM NAME HERE"                    # FILL: full firm name
SALES_REP      = "Sales Rep Name"                    # FILL: rep's full name
OUTPUT_PATH    = "friendly-name/FirmName_Date_Proposal.pptx"  # FILL: output path

# Optional images
WEBSITE_SCREENSHOT_PATH = None   # Set to a local PNG path to show a screenshot; None = skip entirely

# ── Slide 1 ──────────────────────────────────────────────────────
URGENCY_SCORE = "7"                          # FILL: copy from section_05 urgency score

# Pillar status — each is ("RED" | "AMBER" | "GREEN", label, one-line detail)
PILLARS = [
    ("RED",   "CRITICAL", "100% from referrals"),   # FILL: Lead Generation
    ("AMBER", "AMBER",    "Leads falling through"),  # FILL: Intake
    ("AMBER", "AMBER",    "No KPIs / bottleneck"),   # FILL: Team
    ("AMBER", "AMBER",    "No forecasting"),          # FILL: Profit Plan
]
PILLAR_NAMES = ["Lead Generation", "Intake", "Team", "Profit Plan"]

# Key findings — list of ("neg"|"pos", "one-sentence finding")
# Use "neg" for problems (red ✕), "pos" for strengths (green ✓).
# Include 3–4 total; at least one "pos" if a genuine strength exists.
FINDINGS = [
    ("neg", "Finding one — consequence-first, specific to this firm"),
    ("neg", "Finding two — consequence-first, specific to this firm"),
    ("neg", "Finding three — consequence-first, specific to this firm"),
    ("pos", "Finding four — genuine strength of this firm"),
]

# Competitor table — list of (name, "XXX reviews", "brief note")
# List 3 competitors + the client row is added automatically.
COMPETITORS = [
    ("Competitor One",   "XXX reviews", "detail · note"),
    ("Competitor Two",   "XXX reviews", "detail · note"),
    ("Competitor Three", "XXX reviews", "detail · note"),
]
CLIENT_REVIEWS      = "0 reviews"      # FILL: client's actual review count
CLIENT_REVIEWS_NOTE = "← You are here"

# Stage strip (right panel, Slide 1)
STAGE_TEXT = "Stage 4: Small Business Manager  →  Goal: Stage 6, Law Firm Owner"  # FILL

# ── Slide 2 ──────────────────────────────────────────────────────
SLIDE_2_TITLE = "Your Growth Plan: 3 Priorities to Reach $X,XXX/Month"  # FILL

SMB_MODEL_DESC = (
    "All four pillars must work together. "
    "Missing any one means growth stalls regardless of ad spend."
)  # FILL: brief model description

GOAL_HEADLINE = "$X → $Y revenue"       # FILL: e.g. "$1M → $1.8M revenue"
GOAL_DBM      = "Owner takes real time off"  # FILL: DBM outcome phrase

# Each priority: (line1, line2, accent_color_hex, [5 bullet strings])
# Titles must reflect what this firm is actually buying and its real priorities —
# see .claude/commands/audit-pptx.md. Do not default a column to "Marketing Engine"
# unless a marketing/ads package is actually part of this firm's recommendation.
PRIORITIES = [
    (
        "Priority", "One", "0091C9",
        [
            "Bullet 1 — specific action for this firm",
            "Bullet 2 — specific action for this firm",
            "Bullet 3 — specific action for this firm",
            "Bullet 4 — specific action for this firm",
            "Bullet 5 — specific action for this firm",
        ],
    ),
    (
        "Fix Intake &", "Stop Losing Cases", "69CD2B",
        [
            "Bullet 1 — specific action for this firm",
            "Bullet 2 — specific action for this firm",
            "Bullet 3 — specific action for this firm",
            "Bullet 4 — specific action for this firm",
            "Bullet 5 — specific action for this firm",
        ],
    ),
    (
        "Install Team &", "Profit Systems", "003A59",
        [
            "Bullet 1 — specific action for this firm",
            "Bullet 2 — specific action for this firm",
            "Bullet 3 — specific action for this firm",
            "Bullet 4 — specific action for this firm",
            "Bullet 5 — specific action for this firm",
        ],
    ),
]

# ── Slide 3 ──────────────────────────────────────────────────────
# Package cards — (label, bundled_price, retail_price, services_line, accent_color_hex)
PACKAGES = [
    (
        "FULL SERVICE MARKETING — GROWTH",
        "$X,XXX", "$X,XXX/mo",
        "Website · Google Ads · LSA · Meta Ads · GBP optimization",
        "0091C9",
    ),
    (
        "ELITE COACH PLUS",
        "$X,XXX", "$X,XXX/mo",
        "Weekly group coaching · KPI scorecards · Intake framework",
        "003A59",
    ),
]

BUNDLE_TOTAL   = "$X,XXX / mo"                      # FILL: sum of bundled prices
BUNDLE_SAVINGS = "Save $X,XXX/mo by bundling"        # FILL

AD_SPEND_NOTE = (
    "+ Recommended ad spend: $X,XXX–$XX,XXX/mo paid directly to Google/Meta"
)  # FILL — set to None if no marketing/ads package is part of this firm's recommendation

AVG_CASE_VALUE     = "$X,XXX"                         # FILL
CONSERVATIVE_LABEL = "Conservative  (X cases/mo):"   # FILL
CONSERVATIVE_RESULT = "$XX,XXX revenue · X.X× ROAS"  # FILL
AGGRESSIVE_LABEL   = "Aggressive  (X cases/mo):"     # FILL
AGGRESSIVE_RESULT  = "$XX,XXX revenue · X.X× ROAS"   # FILL

# Timeline — 5 items: (milestone_label, action_text)
TIMELINE = [
    ("Day 1",   "Action at day 1 — specific to this firm"),
    ("Day 14",  "Action at day 14 — specific to this firm"),
    ("Week 2",  "Action at week 2 — specific to this firm"),
    ("Week 3",  "Action at week 3 — specific to this firm"),
    ("Month 3", "Action at month 3 — specific to this firm"),
]

CLOSING_QUOTE = (
    '"Closing quote tied to this firm\'s DBM — exact transcript words '
    'or a sharp synthesis of the central opportunity."'
)  # FILL

# ═══════════════════════════════════════════════════════════════════
# LAYOUT ENGINE — DO NOT MODIFY BELOW THIS LINE
# ═══════════════════════════════════════════════════════════════════

# ── Color palette ─────────────────────────────────────────────────
def rgb(h): return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

# SMB Team brand colors (see brand book): Deep Wood Blue, Ocean Blue, Lime Green.
# These three anchor the palette and are unchanged from the previous design —
# what changed in this pass is how they're applied (thin accents/pills instead
# of large solid fills) plus several supporting tones below, retuned to softer,
# lower-saturation values matched from the reference redesign.
NAVY        = rgb("003A59")   # Deep Wood Blue — primary brand color
OCEAN_BLUE  = rgb("0091C9")   # brand accent — large numbers, decorative accents
LIME_GREEN  = rgb("69CD2B")   # brand accent — small "pop" text/labels on navy
WHITE       = rgb("FFFFFF")

BG_WASH       = rgb("F4F7F9")   # full-bleed slide background
URGENCY_CARD  = rgb("0B4C6E")   # header-docked urgency widget
URGENCY_SUB   = rgb("9CC0D6")   # urgency widget's "Score out of 10" subtitle

RED         = rgb("B32D2D")
AMBER       = rgb("B5770B")
GREEN       = rgb("3F8F1A")
RED_PILL_BG   = rgb("FBEEEE")
AMBER_PILL_BG = rgb("FDF4E3")
GREEN_PILL_BG = rgb("ECF7E4")

DARK_TEXT   = rgb("1B2A38")   # near-black body/label text on white cards
BODY_TEXT   = rgb("4A5C6B")   # neutral gray-blue for card copy (findings, ad spend note, timeline actions)
SLATE       = rgb("7C8D9B")   # lighter captions (competitor detail, retail price, services line)
LIGHT_BLUE  = rgb("B8CDDD")   # SMB Team Model description text (on navy)
FOOTER_TEXT = rgb("8FB2C6")

ROI_BG      = rgb("F0FAF4")
ROI_GREEN   = rgb("0B6B4A")
ROI_HILIGHT = rgb("E2F3E9")

BUNDLE_SAVINGS_TEXT = rgb("A9D68C")   # lighter green — reads as "money saved"
GOAL_TEXT_DARK      = rgb("143306")   # goal headline, on Lime Green card
GOAL_TEXT_DARKER    = rgb("1F470C")   # goal DBM line, on Lime Green card

STATUS_COLOR    = {"RED": RED, "AMBER": AMBER, "GREEN": GREEN}
STATUS_PILL_BG  = {"RED": RED_PILL_BG, "AMBER": AMBER_PILL_BG, "GREEN": GREEN_PILL_BG}
# Priority/package accent colors cycle through the brand palette. Lime Green
# is light, so anything filled with it uses a dark green text instead of
# WHITE/navy — see ACCENT_TEXT_COLOR below.
PRIORITY_LIGHT = {
    "0091C9": rgb("ECF5FA"),   # light Ocean Blue tint
    "69CD2B": rgb("ECF7E4"),   # light Lime Green tint
    "003A59": rgb("EEF2F5"),   # light Deep Wood Blue tint
}
# Header text color to use for each accent fill — defaults to WHITE elsewhere.
ACCENT_TEXT_COLOR = {"69CD2B": GOAL_TEXT_DARK}

FONT = "Poppins"
LOGO_PATH = os.path.join(os.path.dirname(__file__), "smb_team_logo.png")
# This script is copied into each firm's own folder (a sibling of "Design
# Files/" at the repo root) to run — the logo above sits right next to it
# because it's copied too, but fonts/ is not, so it's addressed relative to
# the repo root instead.
FONT_DIR = os.path.join(os.path.dirname(__file__), "..", "Design Files", "fonts")

# ── Core helpers ──────────────────────────────────────────────────

def emu(*inches):
    return tuple(int(x * 914400) for x in inches)


# Fields with a documented character budget (see .claude/commands/audit-pptx.md)
# get hard-capped where they're placed on the slide — a backstop so oversized
# content can never spill past its box, no matter what the FILL section
# contains. auto-shrink (set in add_text below) is the first line of defense;
# this cap keeps the shrunk font from having to go microscopic.
_TRUNCATED = []  # (label, original_text) pairs — reported just before saving

def cap(text, max_chars, label=""):
    text = "" if text is None else str(text)
    if len(text) <= max_chars:
        return text
    truncated = text[:max_chars - 1].rstrip()
    if " " in truncated:
        truncated = truncated.rsplit(" ", 1)[0]
    truncated = truncated.rstrip(" ,.;:-—–") + "…"
    _TRUNCATED.append((label or text[:24], text))
    return truncated


def add_rect(slide, left, top, w, h, fill=None, line=False):
    from pptx.util import Emu as E
    shape = slide.shapes.add_shape(1, E(int(left*914400)), E(int(top*914400)),
                                   E(int(w*914400)), E(int(h*914400)))
    shape.line.fill.background() if not line else None
    if not line:
        shape.line.color.rgb = WHITE
        shape.line.width = 0
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    return shape


def add_pill(slide, text, left, top, w, h, bg, text_color, size=6.5, bold=True):
    """A small rounded-rectangle chip with centered text — used for status
    badges (RED/AMBER/GREEN) and the header's urgency-score widget."""
    from pptx.util import Emu as E
    shape = slide.shapes.add_shape(5, E(int(left*914400)), E(int(top*914400)),
                                   E(int(w*914400)), E(int(h*914400)))
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    add_text(slide, text, left, top, w, h, size, text_color, bold=bold, align=PP_ALIGN.CENTER)
    return shape


def add_text(slide, text, left, top, w, h, size, color, bold=False,
             align=PP_ALIGN.LEFT, italic=False, wrap=True,
             cap_chars=None, cap_label=""):
    from pptx.util import Emu as E, Pt as P
    if cap_chars is not None:
        text = cap(text, cap_chars, label=cap_label)
    txb = slide.shapes.add_textbox(E(int(left*914400)), E(int(top*914400)),
                                   E(int(w*914400)), E(int(h*914400)))
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = P(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    # Shrink text to fit the box instead of letting it spill past it —
    # a backstop for any content the character caps above don't cover.
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return txb


def add_image_or_placeholder(slide, path, left, top, w, h, label=""):
    from pptx.util import Emu as E
    if path and os.path.isfile(path):
        slide.shapes.add_picture(path, E(int(left*914400)), E(int(top*914400)),
                                 E(int(w*914400)), E(int(h*914400)))
    else:
        ph = add_rect(slide, left, top, w, h, fill=rgb("D1D5DB"))
        if label:
            add_text(slide, label, left+0.05, top+(h/2)-0.12, w-0.1, 0.24,
                     7, SLATE, align=PP_ALIGN.CENTER)


def add_footer(slide, page, total, logo_path):
    add_rect(slide, 0, 5.27, 10, 0.35, fill=NAVY)
    if logo_path and os.path.isfile(logo_path):
        from pptx.util import Emu as E
        slide.shapes.add_picture(logo_path,
                                 E(int(0.30*914400)), E(int(5.33*914400)),
                                 E(int(1.03*914400)), E(int(0.21*914400)))
    add_text(slide, f"CONFIDENTIAL  ·  PREPARED BY {SALES_REP.upper()} | SMB TEAM",
             1.55, 5.35, 5.40, 0.20, 7, FOOTER_TEXT)
    add_text(slide, f"{page} / {total}", 8.40, 5.35, 1.30, 0.20, 7, FOOTER_TEXT,
             align=PP_ALIGN.RIGHT)


# ── Slide 1: Assessment Overview ──────────────────────────────────

def build_slide1(prs):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)

    # Full-bleed background wash — every card below sits on this, not on white
    add_rect(slide, 0, 0, 10, 5.625, fill=BG_WASH)

    # Header banner + accent stripe (mirrors slides 2 and 3)
    add_rect(slide, 0, 0, 10, 0.92, fill=NAVY)
    add_rect(slide, 0, 0.92, 10, 0.04, fill=LIME_GREEN)
    add_text(slide, f"LAW FIRM GROWTH AUDIT  ·  {FIRM_NAME.upper()}",
             0.32, 0.20, 6.00, 0.18, 7.5, LIME_GREEN, bold=True)
    add_text(slide, f"Where {FIRM_NAME} Stands Today",
             0.32, 0.40, 6.90, 0.40, 20, WHITE, bold=True)

    # Urgency widget — docked in the header, not a large block in the content area
    add_rect(slide, 7.42, 0.18, 2.26, 0.56, fill=URGENCY_CARD)
    add_text(slide, URGENCY_SCORE, 7.58, 0.23, 0.42, 0.42, 21, WHITE, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, "COMPETITIVE URGENCY", 8.04, 0.30, 1.55, 0.18, 7, LIME_GREEN, bold=True)
    add_text(slide, "Score out of 10", 8.04, 0.47, 1.55, 0.18, 7, URGENCY_SUB)

    # Pillar cards — a full-width row of 4, each a white card with a colored
    # top accent bar and a pill-shaped status badge
    pillar_xs = [0.30, 2.68, 5.07, 7.45]
    for i, (status, label, detail) in enumerate(PILLARS):
        x = pillar_xs[i]
        sc = STATUS_COLOR[status]
        add_rect(slide, x, 1.10, 2.26, 0.86, fill=WHITE)
        add_rect(slide, x, 1.10, 2.26, 0.04, fill=sc)
        add_text(slide, PILLAR_NAMES[i], x+0.16, 1.25, 1.94, 0.20, 10, DARK_TEXT, bold=True)
        add_pill(slide, label, x+0.16, 1.50, 0.56, 0.15, STATUS_PILL_BG[status], sc, size=6.5)
        add_text(slide, detail, x+0.16, 1.72, 1.94, 0.20, 8, SLATE,
                 cap_chars=40, cap_label=f"PILLARS[{i}] detail")

    # Key findings label
    add_text(slide, "KEY FINDINGS", 0.30, 2.20, 4.70, 0.18, 7.5, NAVY, bold=True)

    # Finding cards — white with a colored left accent bar and a rounded icon chip
    finding_ys = [2.48, 3.16, 3.84, 4.52]
    for i, (ftype, text) in enumerate(FINDINGS[:4]):
        y = finding_ys[i]
        dot  = RED if ftype == "neg" else GREEN
        sym  = "✕" if ftype == "neg" else "✓"
        add_rect(slide, 0.30, y, 4.72, 0.61, fill=WHITE)
        add_rect(slide, 0.30, y, 0.04, 0.61, fill=dot)
        add_pill(slide, sym, 0.47, y+0.21, 0.18, 0.18, dot, WHITE, size=8)
        add_text(slide, text, 0.75, y+0.13, 3.77, 0.41, 8.5, BODY_TEXT,
                 cap_chars=115, cap_label=f"FINDINGS[{i}]")

    # Website screenshot — only shown when a local PNG is provided
    if WEBSITE_SCREENSHOT_PATH and os.path.isfile(WEBSITE_SCREENSHOT_PATH):
        from pptx.util import Emu as E
        slide.shapes.add_picture(WEBSITE_SCREENSHOT_PATH,
                                 E(int(5.28*914400)), E(int(1.10*914400)),
                                 E(int(4.42*914400)), E(int(1.00*914400)))

    # "You are here" card
    add_rect(slide, 5.28, 2.20, 4.42, 0.74, fill=NAVY)
    add_rect(slide, 5.28, 2.20, 0.04, 0.74, fill=LIME_GREEN)
    add_text(slide, "YOU ARE HERE", 5.48, 2.33, 4.02, 0.16, 7.5, LIME_GREEN, bold=True)
    add_text(slide, STAGE_TEXT, 5.48, 2.54, 4.02, 0.34, 10, WHITE, bold=True)

    # Competitor table header
    add_text(slide, "COMPETITOR LANDSCAPE", 5.28, 3.14, 4.42, 0.18, 7.5, NAVY, bold=True)

    comp_ys = [3.42, 3.86, 4.30]
    for i, (name, reviews, detail) in enumerate(COMPETITORS[:3]):
        y = comp_ys[i]
        add_rect(slide, 5.28, y, 4.42, 0.39, fill=WHITE)
        add_text(slide, name, 5.46, y, 1.78, 0.39, 8.5, DARK_TEXT,
                 cap_chars=30, cap_label=f"COMPETITORS[{i}] name")
        add_text(slide, reviews, 7.26, y, 1.48, 0.39, 8, GREEN,
                 cap_chars=30, cap_label=f"COMPETITORS[{i}] reviews")
        add_text(slide, detail, 8.78, y, 0.80, 0.39, 7, SLATE,
                 cap_chars=34, cap_label=f"COMPETITORS[{i}] detail")

    # Client row — a "you are here" highlight, not a warning, so it uses the
    # same red accent-bar language as the RED pillar/finding cards rather than
    # a solid warning block (which reads as an error next to the neutral rows above it).
    add_rect(slide, 5.28, 4.74, 4.42, 0.39, fill=RED_PILL_BG)
    add_rect(slide, 5.28, 4.74, 0.04, 0.39, fill=RED)
    add_text(slide, FIRM_NAME, 5.46, 4.74, 1.78, 0.39, 8.5, RED, bold=True,
             cap_chars=30, cap_label="CLIENT row firm name")
    add_text(slide, CLIENT_REVIEWS, 7.26, 4.74, 1.48, 0.39, 8, RED, bold=True,
             cap_chars=30, cap_label="CLIENT_REVIEWS")
    add_text(slide, CLIENT_REVIEWS_NOTE, 8.78, 4.74, 0.80, 0.39, 7, SLATE,
             cap_chars=34, cap_label="CLIENT_REVIEWS_NOTE")

    add_footer(slide, 1, 3, LOGO_PATH)


# ── Slide 2: Growth Plan ───────────────────────────────────────────

def build_slide2(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)

    # Background wash + header (mirrors slides 1 and 3)
    add_rect(slide, 0, 0, 10, 5.625, fill=BG_WASH)
    add_rect(slide, 0, 0, 10, 0.92, fill=NAVY)
    add_rect(slide, 0, 0.92, 10, 0.04, fill=LIME_GREEN)
    add_text(slide, f"LAW FIRM GROWTH AUDIT  ·  {FIRM_NAME.upper()}",
             0.32, 0.20, 6.00, 0.18, 7.5, LIME_GREEN, bold=True)
    add_text(slide, SLIDE_2_TITLE, 0.32, 0.40, 9.30, 0.40, 20, WHITE, bold=True)

    # Left column — model card and goal card, as two separate cards with a
    # gap between them rather than one continuous box with an inset panel
    add_rect(slide, 0.30, 1.10, 2.42, 2.50, fill=NAVY)
    add_rect(slide, 0.30, 1.10, 2.42, 0.04, fill=LIME_GREEN)
    add_text(slide, "THE SMB TEAM MODEL", 0.50, 1.28, 2.02, 0.16, 7.5, LIME_GREEN, bold=True)
    add_text(slide, SMB_MODEL_DESC, 0.50, 1.56, 2.02, 1.90, 8.5, LIGHT_BLUE)

    add_rect(slide, 0.30, 3.78, 2.42, 1.23, fill=LIME_GREEN)
    add_text(slide, GOAL_HEADLINE, 0.50, 3.96, 2.02, 0.52, 12, GOAL_TEXT_DARK, bold=True)
    add_text(slide, GOAL_DBM, 0.50, 4.52, 2.02, 0.34, 8.5, GOAL_TEXT_DARKER)

    # 3 priority columns
    priority_xs = [2.92, 5.30, 7.68]
    bullet_ys   = [2.04, 2.65, 3.25, 3.85, 4.46]

    for col_i, (line1, line2, hex_color, bullets) in enumerate(PRIORITIES):
        x = priority_xs[col_i]
        ac = rgb(hex_color)
        light = PRIORITY_LIGHT.get(hex_color, rgb("F8F8FF"))
        header_text = ACCENT_TEXT_COLOR.get(hex_color, WHITE)

        # Header
        add_rect(slide, x, 1.10, 2.26, 0.86, fill=ac)
        add_text(slide, f"0{col_i+1}", x+0.18, 1.22, 0.50, 0.22, 11, header_text, bold=True)
        add_text(slide, line1, x+0.18, 1.46, 1.90, 0.22, 11.5, header_text, bold=True)
        add_text(slide, line2, x+0.18, 1.68, 1.90, 0.22, 11.5, header_text, bold=True)

        # Bullet rows
        for row_i, bullet in enumerate(bullets[:5]):
            y = bullet_ys[row_i]
            bg = light if row_i % 2 == 0 else WHITE
            add_rect(slide, x, y, 2.26, 0.56, fill=bg)
            add_rect(slide, x+0.16, y+0.25, 0.07, 0.07, fill=ac)
            add_text(slide, bullet, x+0.33, y+0.11, 1.76, 0.40, 8, DARK_TEXT,
                     cap_chars=60, cap_label=f"PRIORITIES[{col_i}] bullet {row_i+1}")

    add_footer(slide, 2, 3, LOGO_PATH)


# ── Slide 3: Investment & Next Steps ──────────────────────────────

def build_slide3(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)

    # Background wash + header (mirrors slides 1 and 2)
    add_rect(slide, 0, 0, 10, 5.625, fill=BG_WASH)
    add_rect(slide, 0, 0, 10, 0.92, fill=NAVY)
    add_rect(slide, 0, 0.92, 10, 0.04, fill=LIME_GREEN)
    add_text(slide, f"LAW FIRM GROWTH AUDIT  ·  {FIRM_NAME.upper()}",
             0.32, 0.20, 6.00, 0.18, 7.5, LIME_GREEN, bold=True)
    add_text(slide, "Your Investment & What Happens Next",
             0.32, 0.40, 9.30, 0.40, 20, WHITE, bold=True)

    # Package cards — white with a thin colored left accent bar
    pkg_ys = [1.10, 2.25]
    for i, (label, price, retail, services, hex_c) in enumerate(PACKAGES[:2]):
        y = pkg_ys[i]
        ac = rgb(hex_c)
        add_rect(slide, 0.30, y, 4.42, 1.05, fill=WHITE)
        add_rect(slide, 0.30, y, 0.04, 1.05, fill=ac)
        add_text(slide, label, 0.56, y+0.15, 3.92, 0.16, 7.5, ac, bold=True)
        add_text(slide, price, 0.56, y+0.38, 1.55, 0.42, 25, NAVY, bold=True)
        add_text(slide, "/mo", 1.92, y+0.56, 0.40, 0.20, 10, SLATE)
        add_text(slide, retail, 2.40, y+0.56, 1.00, 0.20, 10, SLATE)
        add_text(slide, services, 0.56, y+0.82, 3.92, 0.18, 8, SLATE,
                 cap_chars=61, cap_label=f"PACKAGES[{i}] services")

    # Bundle total
    add_rect(slide, 0.30, 3.36, 4.42, 0.72, fill=NAVY)
    add_rect(slide, 0.30, 3.36, 0.04, 0.72, fill=LIME_GREEN)
    add_text(slide, "BUNDLE TOTAL", 0.56, 3.51, 1.70, 0.16, 7.5, LIME_GREEN, bold=True)
    add_text(slide, BUNDLE_TOTAL, 0.56, 3.71, 2.20, 0.30, 17, WHITE, bold=True)
    add_text(slide, BUNDLE_SAVINGS, 2.90, 3.78, 1.60, 0.20, 8.5, BUNDLE_SAVINGS_TEXT,
             align=PP_ALIGN.RIGHT)

    # Ad spend note + ROI-on-ad-spend card — both omitted entirely when no
    # marketing/ads package was sold. Showing a return-on-ad-spend projection
    # for ad spend that isn't part of the recommendation is the same bug as
    # the missing-rationale note, just in a second spot on this slide.
    has_ad_spend = bool(AD_SPEND_NOTE)

    if has_ad_spend:
        add_rect(slide, 0.30, 4.16, 4.42, 0.42, fill=rgb("EEF2F5"))
        add_text(slide, AD_SPEND_NOTE, 0.50, 4.16, 4.02, 0.42, 8, BODY_TEXT,
                 cap_chars=120, cap_label="AD_SPEND_NOTE")

        # ROI card
        add_rect(slide, 5.00, 1.10, 4.70, 1.42, fill=ROI_BG)
        add_text(slide, "PROJECTED RETURN ON AD SPEND",
                 5.22, 1.26, 4.26, 0.16, 7.5, ROI_GREEN, bold=True)
        add_text(slide, "Average case value:", 5.22, 1.58, 1.95, 0.18, 8.5, DARK_TEXT, bold=True)
        add_text(slide, AVG_CASE_VALUE, 7.25, 1.58, 2.30, 0.18, 8.5, ROI_GREEN)
        add_rect(slide, 5.14, 1.83, 4.42, 0.28, fill=ROI_HILIGHT)
        add_text(slide, CONSERVATIVE_LABEL, 5.22, 1.90, 1.95, 0.18, 8.5, DARK_TEXT, bold=True)
        add_text(slide, CONSERVATIVE_RESULT, 7.25, 1.90, 2.30, 0.18, 8.5, ROI_GREEN, bold=True)
        add_text(slide, AGGRESSIVE_LABEL, 5.22, 2.21, 1.95, 0.18, 8.5, DARK_TEXT, bold=True)
        add_text(slide, AGGRESSIVE_RESULT, 7.25, 2.21, 2.30, 0.18, 8.5, ROI_GREEN, bold=True)

        first90_y = 2.72
        timeline_ys = [2.99, 3.31, 3.62, 3.94, 4.25]
    else:
        # No ROI card above it — give First 90 Days the full right-column height.
        first90_y = 1.10
        timeline_ys = [1.46, 1.97, 2.48, 2.99, 3.50]

    # First 90 days — alternating light/white rows; the Day-1 dot pops in
    # lime green, the rest in ocean blue, so the immediate next step stands out
    add_text(slide, "WHAT HAPPENS IN THE FIRST 90 DAYS",
             5.00, first90_y, 4.70, 0.18, 7.5, NAVY, bold=True)
    for i, (milestone, action) in enumerate(TIMELINE[:5]):
        y = timeline_ys[i]
        row_bg = rgb("EEF2F5") if i % 2 == 0 else WHITE
        dot_color = LIME_GREEN if i == 0 else OCEAN_BLUE
        add_rect(slide, 5.00, y, 4.70, 0.29, fill=row_bg)
        add_rect(slide, 5.14, y+0.09, 0.10, 0.10, fill=dot_color)
        add_text(slide, milestone, 5.36, y+0.055, 0.75, 0.18, 8, NAVY, bold=True)
        add_text(slide, action, 6.18, y+0.055, 3.35, 0.18, 8, BODY_TEXT,
                 cap_chars=58, cap_label=f"TIMELINE[{i}] action")

    # Closing quote bar — taller, darker navy, with the same lime accent stripe
    add_rect(slide, 0, 4.66, 10, 0.61, fill=rgb("002A41"))
    add_rect(slide, 0, 4.66, 10, 0.02, fill=LIME_GREEN)
    add_text(slide, CLOSING_QUOTE, 0.70, 4.82, 8.60, 0.34, 9.5, rgb("C2D8E6"),
             align=PP_ALIGN.CENTER, cap_chars=200, cap_label="CLOSING_QUOTE")

    add_footer(slide, 3, 3, LOGO_PATH)


# ── Font embedding ────────────────────────────────────────────────
# python-pptx has no font-embedding API, so this patches the saved .pptx's
# raw OOXML directly. Mirrors what sales_companion_template.py already does
# for the PDF via reportlab's registerFont — without it, "Poppins" is just a
# name any viewer without the font installed (Google Slides, a PowerPoint
# missing the font) will silently substitute, which is why fonts can look
# inconsistent across machines even though the script never changes.

def embed_fonts(pptx_path, font_dir):
    import zipfile, shutil, tempfile

    faces = [
        ("regular", "Poppins-Regular.ttf"),
        ("bold",    "Poppins-Bold.ttf"),
        ("italic",  "Poppins-Italic.ttf"),
    ]
    faces = [(role, os.path.join(font_dir, fname)) for role, fname in faces
             if os.path.isfile(os.path.join(font_dir, fname))]
    if not faces:
        return

    NS = {
        "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
        "r":  "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
        "p":  "http://schemas.openxmlformats.org/presentationml/2006/main",
    }

    tmp_dir = tempfile.mkdtemp()
    try:
        with zipfile.ZipFile(pptx_path) as zin:
            zin.extractall(tmp_dir)

        # 1. Copy the font binaries into ppt/fonts/
        fonts_dir = os.path.join(tmp_dir, "ppt", "fonts")
        os.makedirs(fonts_dir, exist_ok=True)
        font_parts = []  # (role, part_filename)
        for i, (role, src_path) in enumerate(faces, start=1):
            part_name = f"font{i}.fntdata"
            shutil.copyfile(src_path, os.path.join(fonts_dir, part_name))
            font_parts.append((role, part_name))

        # 2. Register the .fntdata extension in [Content_Types].xml
        ct_path = os.path.join(tmp_dir, "[Content_Types].xml")
        ct_tree = etree.parse(ct_path)
        ct_root = ct_tree.getroot()
        if not ct_root.xpath("ct:Default[@Extension='fntdata']", namespaces=NS):
            default = etree.SubElement(ct_root, f"{{{NS['ct']}}}Default")
            default.set("Extension", "fntdata")
            default.set("ContentType", "application/x-fontdata")
        ct_tree.write(ct_path, xml_declaration=True, encoding="UTF-8", standalone=True)

        # 3. Add a relationship from presentation.xml to each font part
        rels_path = os.path.join(tmp_dir, "ppt", "_rels", "presentation.xml.rels")
        rels_tree = etree.parse(rels_path)
        rels_root = rels_tree.getroot()
        existing_ids = {el.get("Id") for el in rels_root}

        def new_rid(n=[1]):
            while f"rIdEmbedFont{n[0]}" in existing_ids:
                n[0] += 1
            rid = f"rIdEmbedFont{n[0]}"
            n[0] += 1
            return rid

        rel_ids = {}
        for role, part_name in font_parts:
            rid = new_rid()
            rel = etree.SubElement(rels_root, f"{{{NS['pr']}}}Relationship")
            rel.set("Id", rid)
            rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font")
            rel.set("Target", f"fonts/{part_name}")
            rel_ids[role] = rid
        rels_tree.write(rels_path, xml_declaration=True, encoding="UTF-8", standalone=True)

        # 4. Declare the embedded font in presentation.xml and flip the
        #    embedTrueTypeFonts switch on.
        pres_path = os.path.join(tmp_dir, "ppt", "presentation.xml")
        pres_tree = etree.parse(pres_path)
        pres_root = pres_tree.getroot()
        pres_root.set("embedTrueTypeFonts", "1")

        embedded_font = etree.Element(f"{{{NS['p']}}}embeddedFont")
        font_el = etree.SubElement(embedded_font, f"{{{NS['p']}}}font")
        font_el.set("typeface", FONT)
        for role, rid in rel_ids.items():
            face_el = etree.SubElement(embedded_font, f"{{{NS['p']}}}{role}")
            face_el.set(f"{{{NS['r']}}}id", rid)

        embedded_font_lst = etree.Element(f"{{{NS['p']}}}embeddedFontLst")
        embedded_font_lst.append(embedded_font)

        # Schema order (CT_Presentation) requires embeddedFontLst to sit
        # after notesSz and before defaultTextStyle/custShowLst/etc.
        notes_sz = pres_root.find(f"{{{NS['p']}}}notesSz")
        default_text_style = pres_root.find(f"{{{NS['p']}}}defaultTextStyle")
        if notes_sz is not None:
            insert_at = list(pres_root).index(notes_sz) + 1
        elif default_text_style is not None:
            insert_at = list(pres_root).index(default_text_style)
        else:
            insert_at = len(pres_root)
        pres_root.insert(insert_at, embedded_font_lst)

        pres_tree.write(pres_path, xml_declaration=True, encoding="UTF-8", standalone=True)

        # 5. Re-zip in place
        tmp_pptx = pptx_path + ".tmp"
        with zipfile.ZipFile(tmp_pptx, "w", zipfile.ZIP_DEFLATED) as zout:
            for root_dir, _, files in os.walk(tmp_dir):
                for fname in files:
                    full = os.path.join(root_dir, fname)
                    arcname = os.path.relpath(full, tmp_dir)
                    zout.write(full, arcname)
        shutil.move(tmp_pptx, pptx_path)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


# ── Assemble ──────────────────────────────────────────────────────

prs = Presentation()
prs.slide_width  = Inches(10)
prs.slide_height = Inches(5.625)

build_slide1(prs)
build_slide2(prs)
build_slide3(prs)

os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True) if os.path.dirname(OUTPUT_PATH) else None
prs.save(OUTPUT_PATH)
embed_fonts(OUTPUT_PATH, FONT_DIR)

if _TRUNCATED:
    print(f"WARNING: {len(_TRUNCATED)} field(s) exceeded their character budget "
          f"and were truncated — consider shortening the source copy instead:")
    for label, original in _TRUNCATED:
        print(f"  - {label}: {original!r}")

print(f"Saved: {OUTPUT_PATH}  ({len(prs.slides)} slides)")
